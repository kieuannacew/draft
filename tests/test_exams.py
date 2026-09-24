"""Kiểm tra hàm chấm: file khởi đầu phải SAI hết, file đã làm đúng phải ĐÚNG hết.

Hàm `solve_*` mô phỏng thao tác người học bằng cách ghi ra XML giống Office.
Chạy: python -m pytest -q
"""
import re
import shutil
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table
from pptx import Presentation
from pptx.util import Inches

from mos.core import grade, new_session
from mos.exams import ALL_EXAMS, excel, powerpoint, word


def rewrite_part(path: Path, name: str, fn):
    """Sửa một part XML bên trong file ZIP Office."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        names = zin.namelist()
        for n in names:
            data = zin.read(n)
            if n == name:
                data = fn(data.decode("utf-8")).encode("utf-8")
            zout.writestr(n, data)
        if name not in names:
            zout.writestr(name, fn("").encode("utf-8"))
    shutil.move(tmp, path)


# ------------------------------------------------------------------ Excel

def solve_sales(path):
    wb = load_workbook(path)
    ws = wb["Sheet1"]
    ws.title = "DoanhSo"
    for r in range(2, 12):
        ws[f"E{r}"] = f"=C{r}*D{r}"
        ws[f"D{r}"].number_format = ws[f"E{r}"].number_format = "#,##0"
    ws.add_table(Table(displayName="Table1", ref="A1:E11"))
    ws.freeze_panes = "A2"
    ws["G1"] = "Tổng doanh thu"
    ws["G2"] = "=SUM(E2:E11)"
    wb.save(path)


def solve_staff(path):
    wb = load_workbook(path)
    ws = wb["NhanVien"]
    rows = sorted(excel.STAFF, key=lambda x: -x[2])
    for r, (name, dept, salary) in enumerate(rows, start=2):
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = name, dept, salary
        ws[f"D{r}"] = f'=IF(C{r}>=10000000,"Cao","Thấp")'
    ws.conditional_formatting.add("C2:C11", CellIsRule(operator="greaterThan", formula=["10000000"]))
    wb.defined_names["LuongNV"] = DefinedName("LuongNV", attr_text="NhanVien!$C$2:$C$11")
    ws["F2"] = '=COUNTIF(B2:B11,"Kế toán")'
    chart = BarChart()
    chart.type = "col"
    chart.add_data(Reference(ws, min_col=3, min_row=1, max_row=11), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=11))
    ws.add_chart(chart, "H2")
    ws.page_setup.orientation = "landscape"
    wb.save(path)


# ------------------------------------------------------------------- Word

def _field(instr):
    return parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>'), parse_xml(f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> {instr} </w:instrText></w:r>'), \
        parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>')


def solve_report(path):
    doc = Document(path)
    for p in doc.paragraphs:
        if p.text == "Báo cáo thường niên":
            p.style = doc.styles["Title"]
        elif p.text in word.HEADINGS:
            p.style = doc.styles["Heading 1"]
        for run in p.runs:
            run.text = run.text.replace("ABC", "XYZ")
    toc = doc.paragraphs[0].insert_paragraph_before()
    for r in _field('TOC \\o "1-3" \\h \\z \\u'):
        toc._p.append(r)
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    fp = sec.footer.paragraphs[0]
    for r in _field("PAGE   \\* MERGEFORMAT"):
        fp._p.append(r)
    hp = sec.header.paragraphs[0]
    hp._p.append(parse_xml(
        f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml"><w:pict>'
        '<v:shape id="PowerPlusWaterMarkObject"><v:textpath string="BẢO MẬT"/></v:shape>'
        '</w:pict></w:r>'))
    doc.save(path)


def solve_invite(path):
    doc = Document(path)
    paras = doc.paragraphs
    paras[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for p in paras:
        if p.text in word.BULLETS:
            p._p.get_or_add_pPr().append(parse_xml(
                f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>'))
    # Convert Text to Table: thay 4 dòng tab bằng bảng
    tab_paras = [p for p in paras if "\t" in p.text]
    table = doc.add_table(rows=4, cols=3)
    for i, p in enumerate(tab_paras):
        for j, val in enumerate(p.text.split("\t")):
            table.cell(i, j).text = val
    tab_paras[0]._p.addprevious(table._tbl)
    for p in tab_paras:
        p._p.getparent().remove(p._p)
    doc.core_properties.title = "Thư mời hội thảo"
    doc.save(path)
    rewrite_part(path, "word/settings.xml",
                 lambda x: re.sub(r"(<w:settings[^>]*>)", r"\1<w:trackRevisions/>", x, count=1))
    rewrite_part(path, "word/footnotes.xml", lambda _: (
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:footnote w:id="1"><w:p><w:r><w:t>Danh sách đối tác </w:t></w:r>'
        '<w:r><w:t>đính kèm</w:t></w:r></w:p></w:footnote></w:footnotes>'))


# ------------------------------------------------------------- PowerPoint

P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _sp(xml):
    from pptx.oxml import parse_xml as px
    return px(xml)


def solve_product(path):
    prs = Presentation(path)
    s1 = prs.slides[0]
    s1.notes_slide.notes_text_frame.text = "Chào mừng khán giả"
    new = prs.slides.add_slide(prs.slide_layouts[1])
    new.shapes.title.text = "Kế hoạch"
    for s in prs.slides:
        s._element.append(_sp(f'<p:transition xmlns:p="{P}"><p:fade/></p:transition>'))
    s1._element.append(_sp(
        f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" presetClass="entr"/></p:par>'
        '</p:tnLst></p:timing>'))
    powerpoint._slide(prs, "Giá bán")._element.set("show", "0")
    prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    prs.save(path)


def solve_quarter(path):
    prs = Presentation(path)
    s2 = powerpoint._slide(prs, "Doanh thu theo khu vực")
    s2.shapes.add_table(4, 3, Inches(1), Inches(2), Inches(6), Inches(2))
    s3 = powerpoint._slide(prs, "Quy trình bán hàng")
    s3.shapes._spTree.append(_sp(
        f'<p:graphicFrame xmlns:p="{P}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:nvGraphicFramePr><p:cNvPr id="99" name="Diagram"/><p:cNvGraphicFramePr/><p:nvPr/>'
        '</p:nvGraphicFramePr><p:xfrm><a:off x="0" y="0"/><a:ext cx="1" cy="1"/></p:xfrm>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram"/>'
        '</a:graphic></p:graphicFrame>'))
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        if powerpoint._title(prs.slides.get(int(sld_id.get("id")))) == "Xóa slide này":
            sld_ids.remove(sld_id)
    for i, s in enumerate(prs.slides):
        if i:
            s.shapes._spTree.append(_sp(
                f'<p:sp xmlns:p="{P}"><p:nvSpPr><p:cNvPr id="{50 + i}" name="Slide Number"/>'
                '<p:cNvSpPr/><p:nvPr><p:ph type="sldNum" sz="quarter" idx="12"/></p:nvPr></p:nvSpPr>'
                '<p:spPr/></p:sp>'))
    prs.save(path)
    rewrite_part(path, "ppt/presentation.xml", lambda x: x.replace(
        "</p:presentation>",
        '<p:extLst><p:ext uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}">'
        '<p14:sectionLst xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main">'
        '<p14:section name="Mở đầu" id="{11111111-1111-1111-1111-111111111111}"><p14:sldIdLst/>'
        '</p14:section></p14:sectionLst></p:ext></p:extLst></p:presentation>'))


SOLUTIONS = {
    "DoanhSo.xlsx": solve_sales,
    "NhanVien.xlsx": solve_staff,
    "BaoCao.docx": solve_report,
    "ThuMoi.docx": solve_invite,
    "GioiThieu.pptx": solve_product,
    "BaoCaoQuy.pptx": solve_quarter,
}


@pytest.mark.parametrize("exam", ALL_EXAMS, ids=lambda e: e.code)
def test_fresh_files_score_zero(exam, tmp_path):
    report = grade(new_session(exam, "testing", tmp_path))
    assert report["correct"] == 0
    assert all(not r.error for r in report["results"]), [r.error for r in report["results"]]


@pytest.mark.parametrize("exam", ALL_EXAMS, ids=lambda e: e.code)
def test_solved_files_score_full(exam, tmp_path):
    session = new_session(exam, "testing", tmp_path)
    for i, project in enumerate(exam.projects):
        SOLUTIONS[project.filename](session.file_of(i))
    report = grade(session)
    wrong = [(r.task, r.error) for r in report["results"] if not r.correct]
    assert not wrong
    assert report["score"] == 1000 and report["passed"]


def test_corrupt_file_counts_as_wrong(tmp_path):
    session = new_session(ALL_EXAMS[1], "testing", tmp_path)
    session.file_of(0).write_bytes(b"not a zip")
    report = grade(session)
    assert report["correct"] == 0
    assert any(r.error for r in report["results"])
