"""Lớp học trực tuyến: cấu hình, hàng đợi, báo cáo (luôn chạy) + kịch bản đầy đủ trên Supabase giả (khi có)."""
import json
import uuid

import pytest

from mos import lop_hoc

from . import supabase_gia


# ================================================================ phần chạy mọi nơi
@pytest.fixture
def tmp_dirs(tmp_path, monkeypatch):
    app = tmp_path / "app"
    home = tmp_path / "home"
    data = tmp_path / "data"
    for d in (app, home, data):
        d.mkdir()
    monkeypatch.setattr(lop_hoc, "app_dir", lambda: app)
    monkeypatch.setattr(lop_hoc, "APP_DIR", home)
    monkeypatch.setattr(lop_hoc, "DATA_DIR", data)
    return app, home, data


def test_cau_hinh_may_chu(tmp_dirs):
    app, home, _ = tmp_dirs
    assert lop_hoc.load_server() is None and lop_hoc.Cloud.from_config() is None
    (app / "cau_hinh.json").write_text(json.dumps({"thu_muc_du_lieu": "X:/MOS"}), encoding="utf-8")
    path = lop_hoc.save_server("abc.supabase.co/", " khoa ")
    assert path == app / "cau_hinh.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    assert data["thu_muc_du_lieu"] == "X:/MOS", "không được xóa cấu hình khác"
    assert lop_hoc.load_server() == {"url": "https://abc.supabase.co", "key": "khoa"}
    assert lop_hoc.server_source() == path
    (app / "cau_hinh.json").write_text(json.dumps({"thu_muc_du_lieu": "X:/MOS"}), encoding="utf-8")
    assert lop_hoc.load_server()["url"] == "https://abc.supabase.co", "bản app mới vẫn nhớ máy chủ"
    lop_hoc.save_server("abc.supabase.co/", " khoa ")
    lop_hoc.clear_server()
    assert lop_hoc.load_server() is None
    assert json.loads(path.read_text(encoding="utf-8")) == {"thu_muc_du_lieu": "X:/MOS"}
    (home / "may_chu.json").write_text(json.dumps({"url": "https://b.co", "key": "k"}), encoding="utf-8")
    assert lop_hoc.load_server()["url"] == "https://b.co"


def test_loi_de_hieu():
    assert str(lop_hoc._friendly(400, {"error_code": "invalid_credentials"})) == "Sai tên đăng nhập hoặc mật khẩu."
    assert "Mã lớp" in str(lop_hoc._friendly(400, {"code": "P0001", "message": "sai_ma_lop"}))
    assert "supabase_lop_hoc.sql" in str(lop_hoc._friendly(404, {"code": "PGRST202", "message": "x"}))
    assert str(lop_hoc._friendly(500, {"message": "boom"})).startswith("Máy chủ báo lỗi")


class FakeCloud:
    def __init__(self, fail_after=None):
        self.results, self.progress, self.fail_after = [], [], fail_after

    def _maybe_fail(self):
        if self.fail_after is not None and len(self.results) + len(self.progress) >= self.fail_after:
            raise lop_hoc.OfflineError("mất mạng")

    def send_result(self, data):
        self._maybe_fail()
        self.results.append(data)

    def send_progress(self, data):
        self._maybe_fail()
        self.progress.append(data)


def test_hang_doi(tmp_dirs):
    entry = {"exam": "Đề 1", "code": "WORD", "mode": "testing", "time": "2026-09-01 08:30", "score": 800,
             "correct": 8, "total": 10, "passed": True, "duration": 900, "wrong": ["P1: câu 3"]}
    lop_hoc.queue_result("hv", entry)
    lop_hoc.queue_result("khac", entry)
    lop_hoc.queue_progress("hv", "bai1", "Bài 1", 3, 6, False)
    lop_hoc.queue_progress("hv", "bai1", "Bài 1", 2, 6, True, "80/100")
    assert lop_hoc.pending("hv") == 2
    cloud = FakeCloud(fail_after=1)
    assert lop_hoc.flush(cloud, "hv") == 1
    assert lop_hoc.pending("hv") == 1 and lop_hoc.pending("khac") == 1
    cloud.fail_after = None
    assert lop_hoc.flush(cloud, "hv") == 1
    assert lop_hoc.pending("hv") == 0 and lop_hoc.pending("khac") == 1
    r = cloud.results[0]
    assert r["score"] == 800 and r["passed"] is True and r["lam_luc"].startswith("2026-09-01T08:30")
    assert cloud.progress == [{"bai": "bai1", "ten_bai": "Bài 1", "xem": 3, "tong": 6, "xong": True,
                               "diem": "80/100"}]


def test_tong_hop_va_xuat_excel(tmp_path):
    from openpyxl import load_workbook
    members = [{"id": "a", "username": "an", "ho_ten": "An"}, {"id": "b", "username": "binh", "ho_ten": "Bình"}]
    results = [{"hoc_vien": "a", "code": "WORD", "score": 700, "passed": True, "lam_luc": "2026-09-02T10:00:00"},
               {"hoc_vien": "a", "code": "WORD", "score": 900, "passed": True, "lam_luc": "2026-09-03T10:00:00"},
               {"hoc_vien": "a", "code": "EXCEL", "score": 500, "passed": False, "lam_luc": "2026-09-01T10:00:00"}]
    progress = [{"hoc_vien": "a", "bai": "x", "xong": True}, {"hoc_vien": "b", "bai": "x", "xong": False}]
    rows = lop_hoc.summarize(members, results, progress)
    a, b = rows
    assert a["lan_lam"] == 3 and a["lan_dat"] == 2 and a["best"] == {"WORD": 900, "EXCEL": 500, "POWERPOINT": None}
    assert a["gan_nhat"].startswith("2026-09-03") and a["bai_xong"] == 1
    assert b["lan_lam"] == 0 and b["bai_da_mo"] == 1 and b["bai_xong"] == 0
    out = tmp_path / "lop.xlsx"
    lop_hoc.export_excel(out, "10A1", rows, results, members)
    wb = load_workbook(out)
    assert wb["Tong hop"]["A3"].value == "an" and wb["Tong hop"]["E3"].value == 900
    assert wb["Chi tiet"].max_row == 4


def test_mat_mang():
    c = lop_hoc.Cloud("http://127.0.0.1:9", "k", timeout=2)
    with pytest.raises(lop_hoc.OfflineError):
        c.login("abc", "123456")


# ================================================================ kịch bản trên Supabase giả
need_server = pytest.mark.skipif(not supabase_gia.available(),
                                 reason="cần MOS_TEST_PG + POSTGREST_BIN (xem tests/supabase_gia.py)")


@pytest.fixture(scope="module")
def sb():
    fake = supabase_gia.FakeSupabase().start()
    yield fake
    fake.stop()


def _client(sb):
    return lop_hoc.Cloud(sb.url, supabase_gia.ANON_KEY)


def _result(ma=None, score=800, code="WORD"):
    return lop_hoc.result_payload({"ma": ma or uuid.uuid4().hex, "exam": "Đề 1", "code": code, "mode": "testing",
                                   "time": "2026-09-01 08:30", "score": score, "correct": 8, "total": 10,
                                   "passed": score >= 700, "wrong": ["P1: câu 3"]})


@need_server
def test_ca_lop_hoc(sb):
    admin, gv, hs1, hs2 = _client(sb), _client(sb), _client(sb), _client(sb)
    assert admin.signup("thaygiao", "matkhau1", "Thầy Quản Trị")["vai_tro"] == "quan_tri"
    assert gv.signup("colan", "matkhau2", "Cô Lan")["vai_tro"] == "hoc_vien"
    hs1.signup("an.nguyen", "123456", "Nguyễn An")
    hs2.signup("binh", "123456", "Trần Bình")
    with pytest.raises(lop_hoc.CloudError, match="đã có người dùng"):
        _client(sb).signup("binh", "123456", "Trùng")
    with pytest.raises(lop_hoc.CloudError, match="ít nhất 6"):
        _client(sb).signup("ngan", "123", "Ngắn")
    with pytest.raises(lop_hoc.CloudError, match="Sai tên đăng nhập"):
        _client(sb).login("binh", "sai-mat-khau")

    # --- vai trò: chỉ quản trị cấp quyền giáo viên
    with pytest.raises(lop_hoc.CloudError, match="không có quyền"):
        hs1.set_role(hs1.user_id, "quan_tri")
    with pytest.raises(lop_hoc.CloudError, match="không có quyền"):
        gv.create_class("Lớp lậu")
    admin.set_role(gv.user_id, "giao_vien")
    assert gv.load_profile()["vai_tro"] == "giao_vien"
    with pytest.raises(lop_hoc.CloudError, match="ít nhất một tài khoản quản trị"):
        admin.set_role(admin.user_id, "hoc_vien")

    # --- tạo lớp, vào lớp bằng mã
    lop = gv.create_class("10A1")
    assert len(lop["ma"]) == 6
    with pytest.raises(lop_hoc.CloudError, match="Mã lớp không đúng"):
        hs1.join_class("ZZZZZZ")
    assert hs1.join_class(lop["ma"].lower())["ten"] == "10A1"
    hs1.join_class(lop["ma"])                                  # vào lại không lỗi
    mine = hs1.my_classes()
    assert [c["ten"] for c in mine] == ["10A1"] and mine[0]["ma"] == "" and mine[0]["ten_giao_vien"] == "Cô Lan"
    assert gv.my_classes()[0]["ma"] == lop["ma"] and gv.my_classes()[0]["si_so"] == 1
    assert hs2.my_classes() == []
    with pytest.raises(lop_hoc.CloudError):
        hs2.join_class("")

    # --- kết quả: gửi 2 lần cùng mã chỉ lưu 1; không gửi thay người khác
    r = _result()
    hs1.send_result(r)
    hs1.send_result(r)
    hs1.send_result(_result(score=950, code="EXCEL"))
    hs2.send_result(_result(score=300))
    assert len(hs1.results()) == 2
    fake = dict(_result(), hoc_vien=hs2.user_id)
    with pytest.raises(lop_hoc.CloudError):
        hs1.send_result(fake)
    assert len(hs2.results()) == 1

    # --- giáo viên chỉ thấy học sinh trong lớp mình
    members = gv.members(lop["id"])
    assert [m["username"] for m in members] == ["an.nguyen"]
    seen = gv.results([hs1.user_id, hs2.user_id])
    assert {x["hoc_vien"] for x in seen} == {hs1.user_id} and len(seen) == 2
    assert {p["username"] for p in gv.profiles()} == {"colan", "an.nguyen"}
    assert len(admin.profiles()) == 4
    assert len(hs1.profiles()) == 1
    assert hs1.members(lop["id"]) == [] or all(m["id"] == hs1.user_id for m in hs1.members(lop["id"]))

    # --- tiến độ bài giảng chỉ tăng
    hs1.send_progress({"bai": "b1", "ten_bai": "Bài 1", "xem": 3, "tong": 6, "xong": False})
    hs1.send_progress({"bai": "b1", "ten_bai": "Bài 1", "xem": 2, "tong": 6, "xong": True, "diem": "80/100"})
    hs1.send_progress({"bai": "b1", "ten_bai": "Bài 1", "xem": 1, "tong": 6, "xong": False})
    p = gv.progress([hs1.user_id])[0]
    assert (p["xem"], p["xong"], p["diem"]) == (3, True, "80/100")
    assert gv.progress([hs2.user_id]) == []

    summary = lop_hoc.summarize(members, seen, gv.progress([m["id"] for m in members]))
    assert summary[0]["best"]["EXCEL"] == 950 and summary[0]["bai_xong"] == 1

    # --- sửa / xóa lớp: chỉ giáo viên của lớp
    hs1.rename_class(lop["id"], "Đổi trộm")                    # RLS: không có dòng nào bị sửa
    assert gv.my_classes()[0]["ten"] == "10A1"
    gv.rename_class(lop["id"], "10A1 – Tin học")
    assert hs1.my_classes()[0]["ten"] == "10A1 – Tin học"

    # --- đặt lại mật khẩu
    gv.reset_password(hs1.user_id, "moi123")
    assert _client(sb).login("an.nguyen", "moi123")["username"] == "an.nguyen"
    with pytest.raises(lop_hoc.CloudError, match="không có quyền"):
        gv.reset_password(hs2.user_id, "moi123")
    with pytest.raises(lop_hoc.CloudError, match="không có quyền"):
        gv.reset_password(admin.user_id, "moi123")
    with pytest.raises(lop_hoc.CloudError, match="ít nhất 6"):
        admin.reset_password(hs2.user_id, "123")

    # --- tự đổi mật khẩu, họ tên
    hs2.change_password("binh2026")
    hs2.rename_me("Trần Văn Bình")
    assert _client(sb).login("binh", "binh2026")["ho_ten"] == "Trần Văn Bình"

    # --- khóa tài khoản
    admin.set_locked(hs2.user_id, True)
    with pytest.raises(lop_hoc.CloudError, match="bị khóa"):
        _client(sb).login("binh", "binh2026")
    with pytest.raises(lop_hoc.CloudError):
        hs2.send_result(_result())                              # phiên cũ cũng không gửi được nữa
    with pytest.raises(lop_hoc.CloudError, match="tự khóa"):
        admin.set_locked(admin.user_id, True)
    admin.set_locked(hs2.user_id, False)

    # --- học sinh rời lớp / giáo viên xóa học sinh khỏi lớp
    hs2.join_class(lop["ma"])
    assert gv.my_classes()[0]["si_so"] == 2
    gv.leave_class(lop["id"], hs2.user_id)
    hs1.leave_class(lop["id"])
    assert gv.members(lop["id"]) == []

    # --- xóa lớp, xóa tài khoản
    hs2.delete_class(lop["id"])                                 # không phải lớp của mình: không xóa được
    assert len(gv.my_classes()) == 1
    gv.delete_class(lop["id"])
    assert gv.my_classes() == []
    with pytest.raises(lop_hoc.CloudError, match="tự xóa"):
        admin.delete_account(admin.user_id)
    admin.delete_account(hs2.user_id)
    with pytest.raises(lop_hoc.CloudError, match="Sai tên đăng nhập"):
        _client(sb).login("binh", "binh2026")


@need_server
def test_kiem_tra_may_chu(sb):
    assert _client(sb).check() == {"schema": True, "autoconfirm": True, "signup": True}
    with pytest.raises(lop_hoc.CloudError, match="anon key"):
        lop_hoc.Cloud(sb.url, "khoa-sai").check()


@need_server
def test_khoa_kieu_moi(sb):
    c = lop_hoc.Cloud(sb.url, supabase_gia.PUBLISHABLE_KEY)
    assert c.check()["schema"] is True
    c.signup("khoa.moi", "123456", "Khóa Mới")
    assert c.my_classes() == []


@need_server
def test_het_han_phien_tu_lam_moi(sb):
    c = _client(sb)
    c.signup("hethan", "123456", "Hết Hạn")
    sb.token_seconds = -10
    try:
        c.login("hethan", "123456")                             # token nhận về đã hết hạn
    finally:
        sb.token_seconds = 3600
    assert c.my_classes() == []                                 # tự làm mới phiên rồi gọi lại


@need_server
def test_gui_hang_doi_len_may_chu(sb, tmp_dirs):
    c = _client(sb)
    c.signup("hangdoi", "123456", "Hàng Đợi")
    entry = {"exam": "Đề 2", "code": "POWERPOINT", "mode": "training", "time": "2026-09-05 14:00", "score": 1000,
             "correct": 5, "total": 5, "passed": True, "wrong": []}
    lop_hoc.queue_result("hangdoi", entry)
    lop_hoc.queue_progress("hangdoi", "bai-x", "Bài X", 1, 1, True)
    assert lop_hoc.flush(c, "hangdoi") == 2 and lop_hoc.pending("hangdoi") == 0
    assert c.results()[0]["score"] == 1000
    assert c.progress()[0]["xong"] is True


# ================================================================ tài khoản bản sao + giao diện
def test_ban_sao_tai_khoan(tmp_path):
    from mos.accounts import AccountError, AccountStore
    store = AccountStore(tmp_path / "tk.json")
    acc = store.mirror("An.Nguyen", "Nguyễn An", "giao_vien", "matkhau1")
    assert acc.nguon == "may_chu" and acc.is_teacher and not acc.is_admin and not acc.doi_mat_khau
    assert store.authenticate("an.nguyen", "matkhau1").ho_ten == "Nguyễn An"
    store.mirror("an.nguyen", "Nguyễn Văn An", "vai_tro_la", "moi456")
    with pytest.raises(AccountError):
        store.authenticate("an.nguyen", "matkhau1")
    acc = store.authenticate("an.nguyen", "moi456")
    assert acc.vai_tro == "hoc_vien" and acc.ho_ten == "Nguyễn Văn An"
    assert AccountStore(tmp_path / "tk.json").get("an.nguyen").nguon == "may_chu"


@pytest.fixture
def sb_moi():
    """Máy chủ giả riêng (CSDL trống) – người đăng ký đầu tiên là quản trị."""
    fake = supabase_gia.FakeSupabase().start()
    yield fake
    fake.stop()


@need_server
def test_giao_dien_lop_hoc(sb_moi, tmp_dirs, monkeypatch):
    sb = sb_moi
    pytest.importorskip("PySide6")
    from PySide6.QtWidgets import QApplication
    from mos.accounts import AccountStore
    from mos.qt import theme as T
    from mos.qt import app as A
    from mos.qt import lop
    QApplication.instance() or QApplication([])
    monkeypatch.setattr(T, "info", lambda *a: None)
    monkeypatch.setattr(T, "error", lambda *a: (_ for _ in ()).throw(AssertionError(a[2])))
    lop_hoc.save_server(sb.url, supabase_gia.ANON_KEY)
    gv = _client(sb)
    assert gv.signup("gv.ui", "123456", "Cô Giao Diện")["vai_tro"] == "quan_tri"
    lop_moi = gv.create_class("12C3")
    store = AccountStore(tmp_dirs[2] / "tk.json")
    win = A.MainWindow(store)

    def login(user, pw):
        page = win.stack.currentWidget()
        assert isinstance(page, A.LoginPage) and page.server
        page.user.setText(user)
        page.pw.setText(pw)
        page.login()
        return win.stack.currentWidget()

    shell = login("gv.ui", "123456")
    assert isinstance(shell, A.Shell) and win.cloud is not None and "lop" in shell.nav_buttons
    shell.go("lop")
    assert isinstance(shell.content.currentWidget(), lop.ClassPage)
    win.logout()

    hs = _client(sb)
    hs.signup("hs.ui", "123456", "Học Sinh UI")
    hs.join_class(lop_moi["ma"])
    shell = login("hs.ui", "123456")
    assert store.get("hs.ui").nguon == "may_chu"
    shell.go("lop")
    shell.sync_lesson(type("L", (), {"id": "b", "ten": "Bài", "count": 3})(), 3, True)
    import time
    for _ in range(50):
        if lop_hoc.pending("hs.ui") == 0:
            break
        time.sleep(0.1)
    assert hs.progress()[0]["xong"] is True
    win.logout()

    shell = login("gv.ui", "123456")
    shell.go("lop")
    page = shell.content.currentWidget()
    assert [m["username"] for m in page.members] == ["hs.ui"]
    win.logout()

    # mất mạng: vẫn đăng nhập bằng bản sao; tài khoản chưa từng đăng nhập thì không
    lop_hoc.save_server("http://127.0.0.1:9", supabase_gia.ANON_KEY)
    win.logout()
    shell = login("hs.ui", "123456")
    assert isinstance(shell, A.Shell) and win.cloud is None
    win.logout()
    page = win.stack.currentWidget()
    page.user.setText("chua.co")
    page.pw.setText("123456")
    page.login()
    assert win.stack.currentWidget() is page and "máy chủ" in page.err.text()
