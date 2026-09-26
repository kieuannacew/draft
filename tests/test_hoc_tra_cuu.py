"""Luyện theo chương, tài liệu học (bài giảng PPTX) và tra từ khóa."""
import os
import random

import pytest

from mos import chuong, tai_lieu, tu_khoa
from mos.custom import load_custom_exams, merge_exams
from mos.exams import ALL_EXAMS

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


@pytest.fixture(scope="module")
def exams():
    custom, _ = load_custom_exams()
    return merge_exams(ALL_EXAMS, custom)


# ---------------------------------------------------------------- chương
def test_moi_mon_co_chuong():
    for code in ("WORD", "EXCEL", "POWERPOINT"):
        assert chuong.chapters(code)
        assert chuong.chapter(code, 1).number == 1
    assert chuong.chapter("WORD", 99) is None


def test_moi_cau_de_co_san_deu_co_chuong():
    for exam in ALL_EXAMS:
        for p in exam.projects:
            for t in p.tasks:
                assert chuong.chapter(exam.code, t.chapter), (exam.name, t.title)


def test_bai_luyen_theo_chuong(exams):
    counts = chuong.count_tasks(exams, "WORD")
    assert counts and all(n > 0 for n in counts.values())
    for number in counts:
        exam = chuong.practice_exam(exams, "WORD", number, rng=random.Random(1))
        assert exam is not None and exam.minutes == 0
        assert exam.name.startswith(f"Chương {number}:")
        names = [p.filename.casefold() for p in exam.projects]
        assert len(names) == len(set(names)), "tên file trùng sẽ ghi đè nhau"
        assert len(exam.projects) <= 6
        for p in exam.projects:
            assert p.tasks and all(t.chapter == number for t in p.tasks)
    assert chuong.practice_exam(exams, "WORD", 99) is None


# ---------------------------------------------------------------- tài liệu học
def _make_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Chèn hình mờ"
    s.placeholders[1].text = "Design → Watermark → Custom Watermark"
    s.notes_slide.notes_text_frame.text = "Ghi chú cho giáo viên"
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Bảng"
    s.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table.cell(0, 0).text = "Ô A1"
    prs.save(path)


def test_dong_goi_anh_khong_doc_truc_tiep(tmp_path):
    imgs = [b"\x89PNG anh 1", b"\x89PNG anh 2"]
    tai_lieu.pack_slides(tmp_path, "bai-1", imgs)
    raw = (tmp_path / tai_lieu.PACK).read_bytes()
    assert b"PNG anh" not in raw
    assert tai_lieu.read_slide(tmp_path, "bai-1", 2) == imgs[1]


def test_nhap_pptx_va_tien_do(tmp_path, monkeypatch):
    pytest.importorskip("PySide6")
    src = tmp_path / "bai.pptx"
    _make_pptx(src)
    base = tmp_path / "tai_lieu"
    lesson = tai_lieu.import_pptx(src, "Hình mờ", "word", chuong=2, dest_base=base, use_powerpoint=False)
    assert lesson.count == 2 and lesson.mon == "WORD" and lesson.chuong == 2
    assert "Watermark" in lesson.slides[0]["text"]
    assert lesson.slides[0]["notes"] == "Ghi chú cho giáo viên"
    assert "Ô A1" in lesson.slides[1]["text"]
    assert lesson.image(1).startswith(b"\x89PNG")
    assert not list(base.rglob("*.pptx")), "không được chép file pptx gốc"
    # tìm thấy trong danh sách, tra cứu được slide
    monkeypatch.setattr(tai_lieu, "search_dirs", lambda: [base])
    assert [x.id for x in tai_lieu.list_lessons()] == [lesson.id]
    hits = tu_khoa.search_slides("watermark", [lesson])
    assert hits and hits[0].number == 1
    # tiến độ chỉ tăng
    monkeypatch.setattr(tai_lieu, "DATA_DIR", tmp_path)
    tai_lieu.save_progress("hv", lesson.id, 2)
    tai_lieu.save_progress("hv", lesson.id, 1)
    assert tai_lieu.load_progress("hv") == {lesson.id: 2}
    assert tai_lieu.load_progress("khac") == {}
    tai_lieu.delete_lesson(lesson)
    assert not lesson.folder.exists()


def test_bai_giang_mau_doc_duoc():
    lessons = [x for x in tai_lieu.list_lessons() if x.folder.parent.name == "tai_lieu"]
    assert {x.loai for x in lessons} >= {"slides", "scorm"}
    for x in lessons:
        assert x.count > 0
        if x.loai == "slides":
            assert x.image(1).startswith(b"\x89PNG")
        if x.loai == "scorm":
            assert x.launch_path.is_file()


def test_ma_hoa_hai_chieu():
    key = tai_lieu._key("x")
    data = bytes(range(256)) * 1000 + b"\x00\x00"
    enc = tai_lieu._xor(data, key)
    assert enc != data and len(enc) == len(data)
    assert tai_lieu._xor(enc, key) == data
    assert tai_lieu._xor(b"", key) == b""


# ---------------------------------------------------------------- video
def test_nhap_video(tmp_path, monkeypatch):
    src = tmp_path / "bai.mp4"
    raw = b"\x00\x00\x00\x18ftypmp42" + bytes(5000)
    src.write_bytes(raw)
    lesson = tai_lieu.import_file(src, "Video hình mờ", "word", 1, mo_ta="chèn watermark",
                                  dest_base=tmp_path / "tl")
    assert lesson.loai == "video" and lesson.count == 1 and lesson.chuong == 1
    assert b"ftyp" not in (lesson.folder / tai_lieu.VIDEO).read_bytes()
    assert lesson.video() == raw
    assert not list((tmp_path / "tl").rglob("*.mp4"))
    assert tu_khoa.search_slides("watermark", [lesson])
    with pytest.raises(ValueError):
        tai_lieu.import_video(tmp_path / "x.txt", "x", "WORD", dest_base=tmp_path / "tl")


# ---------------------------------------------------------------- SCORM
MANIFEST_12 = b"""<?xml version="1.0"?>
<manifest xmlns="http://www.imsproject.org/xsd/imscp_rootv1p1p2"
          xmlns:adlcp="http://www.adlnet.org/xsd/adlcp_rootv1p2">
  <metadata><schemaversion>1.2</schemaversion></metadata>
  <organizations default="o"><organization identifier="o"><title>Khoa hoc</title>
    <item identifier="i" identifierref="r"><title>Bai 1 Freeze Panes</title></item></organization></organizations>
  <resources><resource identifier="r" type="webcontent" adlcp:scormtype="sco" xml:base="noi_dung/" href="index.html?a=1"/>
  </resources></manifest>"""

MANIFEST_2004 = b"""<?xml version="1.0"?>
<manifest xmlns="http://www.imsglobal.org/xsd/imscp_v1p1" xmlns:adlcp="http://www.adlnet.org/xsd/adlcp_v1p3">
  <metadata><schemaversion>2004 4th Edition</schemaversion></metadata>
  <organizations/><resources><resource identifier="r" href="story.html"/></resources></manifest>"""


def test_doc_manifest():
    info = tai_lieu.read_manifest(MANIFEST_12)
    assert info["launch"] == "noi_dung/index.html" and info["phien_ban"] == "1.2"
    assert "Bai 1 Freeze Panes" in info["tieu_de"]
    info = tai_lieu.read_manifest(MANIFEST_2004)
    assert info["launch"] == "story.html" and info["phien_ban"] == "2004"
    with pytest.raises(ValueError):
        tai_lieu.read_manifest(b"<manifest/>")


def test_nhap_scorm_zip(tmp_path):
    import zipfile
    pkg = tmp_path / "goi.zip"
    with zipfile.ZipFile(pkg, "w") as z:
        z.writestr("goi/imsmanifest.xml", MANIFEST_12)
        z.writestr("goi/noi_dung/index.html", "<html>xin chao</html>")
        z.writestr("goi/../../tron.txt", "khong duoc ghi ra ngoai")
    lesson = tai_lieu.import_file(pkg, "Bài tương tác", "EXCEL", 1, dest_base=tmp_path / "tl")
    assert lesson.loai == "scorm" and lesson.phien_ban == "1.2"
    assert lesson.launch_path.read_text() == "<html>xin chao</html>"
    assert not (tmp_path / "tron.txt").exists() and not (tmp_path / "tl" / "tron.txt").exists()
    assert tu_khoa.search_slides("freeze panes", [lesson])
    bad = tmp_path / "khong_phai.zip"
    with zipfile.ZipFile(bad, "w") as z:
        z.writestr("a.txt", "x")
    with pytest.raises(ValueError):
        tai_lieu.import_file(bad, "x", "WORD", dest_base=tmp_path / "tl")
    with pytest.raises(ValueError):
        tai_lieu.import_file(tmp_path / "a.docx", "x", "WORD", dest_base=tmp_path / "tl")


def test_du_lieu_scorm(tmp_path, monkeypatch):
    monkeypatch.setattr(tai_lieu, "DATA_DIR", tmp_path)
    assert tai_lieu.load_scorm("hv", "b") == {}
    tai_lieu.save_scorm("hv", "b", {"cmi.core.lesson_status": "incomplete", "cmi.suspend_data": "3"})
    assert tai_lieu.load_progress("hv") == {}
    assert tai_lieu.load_scorm("hv", "b")["cmi.suspend_data"] == "3"
    cmi = {"cmi.core.lesson_status": "passed", "cmi.core.score.raw": "80", "cmi.core.score.max": "100"}
    tai_lieu.save_scorm("hv", "b", cmi)
    assert tai_lieu.load_progress("hv") == {"b": 1}
    assert tai_lieu.scorm_score(cmi) == "80/100"
    assert tai_lieu.scorm_done({"cmi.completion_status": "completed"})
    assert tai_lieu.scorm_score({"cmi.score.scaled": "0.75"}) == "75%"
    assert tai_lieu.scorm_score({}) == ""


def test_api_scorm_trong_trang():
    pytest.importorskip("PySide6")
    from mos.qt import hoc
    js = hoc.scorm_script({"cmi.suspend_data": "2"}, "hv")
    assert "window.API =" in js and "window.API_1484_11" in js and '"cmi.suspend_data": "2"' in js
    assert "__DATA__" not in js and "__PREFIX__" not in js
    assert hoc.parse_scorm_message('MOS_SCORM:commit:{"cmi.core.lesson_status": "passed"}') == {
        "cmi.core.lesson_status": "passed"}
    assert hoc.parse_scorm_message("log khac") is None
    assert hoc.parse_scorm_message("MOS_SCORM:set:khong-phai-json") is None


# ---------------------------------------------------------------- tra từ khóa
def test_fold_bo_dau():
    assert tu_khoa.fold("Mục Lục") == tu_khoa.fold("muc luc")
    assert tu_khoa.fold("Đường viền") == "duong vien"


@pytest.mark.parametrize("q", ["watermark", "hình mờ", "hinh mo", "Freeze Panes", "mục lục"])
def test_tra_thuat_ngu(q):
    assert tu_khoa.search_terms(q)


def test_tra_theo_mon():
    terms = tu_khoa.search_terms("freeze", "EXCEL")
    assert terms and all(t.subject in ("EXCEL", "ALL") for t in terms)
    assert not tu_khoa.search_terms("freeze", "POWERPOINT")


def test_tra_cau_trong_de(exams):
    hits = tu_khoa.search_tasks("watermark", exams)
    assert hits
    assert all(h.exam.code == "WORD" for h in hits)
    assert tu_khoa.search_tasks("zzzkhongcozzz", exams) == []


def test_anh_slide_theo_mat_do_diem_anh():
    """Slide co theo mật độ điểm ảnh thật (Windows phóng to 150% vẫn nét)."""
    pytest.importorskip("PySide6")
    from PySide6.QtGui import QPixmap
    from PySide6.QtWidgets import QApplication, QLabel
    from mos.qt import hoc
    QApplication.instance() or QApplication([])
    w = QLabel()
    src = QPixmap(2560, 1440)
    out = hoc.sharp(src, 800, 600, w)
    dpr = w.devicePixelRatioF()
    assert out.devicePixelRatio() == dpr
    assert out.width() == round(800 * dpr) and out.height() == round(450 * dpr)
    assert tai_lieu.SLIDE_WIDTH >= 2560
