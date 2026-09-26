"""Tài liệu học (bài giảng PowerPoint) xem ngay trong app, không cho tải về.

Mỗi bài giảng là một thư mục `tai_lieu/<id>/`, có 3 loại ("loai" trong bai.json):
    slides  bai.json (tên Việt/Anh, môn, chương, chữ + ghi chú từng slide để tra cứu)
            + slides.mosl: ảnh các slide, đóng gói và mã hóa (không mở được bằng trình xem ảnh / PowerPoint)
    video   bai.json + video.mosv: video ngắn (mp4…) đã mã hóa, chỉ giải mã trong bộ nhớ khi phát
    scorm   bai.json + scorm/: gói SCORM 1.2 / 2004 (HTML tương tác) giải nén, "launch" = trang mở đầu

Nhập bài từ file .pptx:
  1. Trên Windows có PowerPoint: PowerPoint xuất từng slide thành ảnh PNG (đẹp y như bản gốc).
  2. Không có PowerPoint: app tự dựng ảnh slide từ chữ, hình và ảnh trong file (đơn giản hơn).
File .pptx gốc KHÔNG được chép vào app. Học viên chỉ xem ảnh trong trình xem, không có nút lưu / xuất.
(Không chặn được chụp màn hình – điều này đúng với mọi phần mềm.)
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from .core import DATA_DIR, app_dir

PACK = "slides.mosl"
VIDEO = "video.mosv"
SCORM_DIR = "scorm"
VIDEO_EXT = (".mp4", ".m4v", ".mov", ".webm", ".mkv", ".avi", ".wmv")
META = "bai.json"
SLIDE_WIDTH = 2560          # đủ nét cho màn hình 2K và Windows phóng to 150%
_MAGIC = b"MOSL1"
_QT_APP = None


def search_dirs() -> list[Path]:
    return [app_dir() / "tai_lieu", DATA_DIR / "tai_lieu"]


def editable_dir() -> Path:
    return DATA_DIR / "tai_lieu"


# ====================================================================== đóng gói ảnh


def _key(lesson_id: str) -> bytes:
    return hashlib.sha256(("mos-bai-giang:" + lesson_id).encode("utf-8")).digest()


def _xor(data: bytes, key: bytes) -> bytes:
    if not data:
        return b""
    k = (key * (len(data) // len(key) + 1))[:len(data)]
    return (int.from_bytes(data, "little") ^ int.from_bytes(k, "little")).to_bytes(len(data), "little")


def pack_slides(folder: Path, lesson_id: str, images: list[bytes]) -> None:
    key = _key(lesson_id)
    with zipfile.ZipFile(folder / PACK, "w", zipfile.ZIP_STORED) as z:
        z.writestr("id", _MAGIC + lesson_id.encode("utf-8"))
        for i, img in enumerate(images, start=1):
            z.writestr(f"{i:04d}", _xor(img, key))


def read_slide(folder: Path, lesson_id: str, number: int) -> bytes:
    """Ảnh PNG của slide thứ `number` (1, 2, …) – chỉ giải mã trong bộ nhớ."""
    with zipfile.ZipFile(folder / PACK) as z:
        return _xor(z.read(f"{number:04d}"), _key(lesson_id))


# ====================================================================== dữ liệu bài giảng


@dataclass
class Lesson:
    id: str
    folder: Path
    ten: str
    ten_en: str = ""
    mon: str = "WORD"
    chuong: int | None = None
    mo_ta: str = ""
    mo_ta_en: str = ""
    slides: list[dict] = field(default_factory=list)   # [{"text": ..., "notes": ...}]
    loai: str = "slides"                                # slides / video / scorm
    launch: str = ""                                    # scorm: trang mở đầu (tương đối trong scorm/)
    phien_ban: str = ""                                 # scorm: "1.2" / "2004"

    @property
    def count(self) -> int:
        """Số bước để tính tiến độ: slide = số slide; video / SCORM = 1 (xem hết / hoàn thành)."""
        return len(self.slides) if self.loai == "slides" else 1

    @property
    def launch_path(self) -> Path:
        return self.folder / SCORM_DIR / self.launch

    def video(self) -> bytes:
        """Dữ liệu video (đã giải mã, chỉ trong bộ nhớ)."""
        return _xor((self.folder / VIDEO).read_bytes(), _key(self.id))

    def image(self, number: int) -> bytes:
        return read_slide(self.folder, self.id, number)


def load_lesson(folder: Path) -> Lesson | None:
    try:
        data = json.loads((folder / META).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    loai = data.get("loai", "slides")
    need = {"slides": folder / PACK, "video": folder / VIDEO,
            "scorm": folder / SCORM_DIR / str(data.get("launch", ""))}.get(loai)
    if need is None or not need.is_file():
        return None
    return Lesson(id=data.get("id", folder.name), folder=folder, ten=data.get("ten", folder.name),
                  ten_en=data.get("ten_en", ""), mon=str(data.get("mon", "WORD")).upper(),
                  chuong=data.get("chuong"), mo_ta=data.get("mo_ta", ""), mo_ta_en=data.get("mo_ta_en", ""),
                  slides=list(data.get("slides", [])), loai=loai, launch=str(data.get("launch", "")),
                  phien_ban=str(data.get("phien_ban", "")))


def list_lessons() -> list[Lesson]:
    out, seen = [], set()
    for base in search_dirs():
        if not base.is_dir():
            continue
        for folder in sorted(p for p in base.iterdir() if p.is_dir()):
            lesson = load_lesson(folder)
            if lesson and lesson.id not in seen:
                seen.add(lesson.id)
                out.append(lesson)
    order = {"WORD": 0, "EXCEL": 1, "POWERPOINT": 2}
    return sorted(out, key=lambda x: (order.get(x.mon, 9), x.chuong or 99, x.ten.casefold()))


def delete_lesson(lesson: Lesson) -> None:
    shutil.rmtree(lesson.folder, ignore_errors=True)


def _slug(text: str) -> str:
    text = unicodedata.normalize("NFKD", text.replace("đ", "d").replace("Đ", "D"))
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_").lower() or "bai"


# ====================================================================== đọc chữ trong pptx


def _shape_text(shape) -> list[str]:
    out = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        out += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            out.append("  |  ".join(c.text for c in row.cells))
    if shape.shape_type == 6:           # nhóm
        for s in shape.shapes:
            out += _shape_text(s)
    return out


def extract_text(pptx: Path) -> list[dict]:
    from pptx import Presentation
    prs = Presentation(str(pptx))
    slides = []
    for slide in prs.slides:
        lines = []
        for shape in slide.shapes:
            lines += _shape_text(shape)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append({"text": "\n".join(lines), "notes": notes})
    return slides


# ====================================================================== ảnh slide


def export_with_powerpoint(pptx: Path, out_dir: Path, width: int = SLIDE_WIDTH) -> list[bytes]:
    """Nhờ PowerPoint (Windows) xuất mỗi slide thành PNG. Lỗi / không có PowerPoint → []."""
    if os.name != "nt":
        return []
    script = (
        "$ErrorActionPreference='Stop';"
        "$pp = New-Object -ComObject PowerPoint.Application;"
        f"$p = $pp.Presentations.Open('{pptx}', $true, $false, $false);"
        f"$h = [int]({width} * $p.PageSetup.SlideHeight / $p.PageSetup.SlideWidth);"
        f"$p.Export('{out_dir}', 'PNG', {width}, $h);"
        "$p.Close();"
        "if ($pp.Presentations.Count -eq 0) { $pp.Quit() }"
    )
    try:
        subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                       check=True, capture_output=True, timeout=300,
                       creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    except (OSError, subprocess.SubprocessError):
        return []
    files = [f for f in out_dir.glob("*") if f.suffix.lower() == ".png"]
    files.sort(key=lambda f: int(re.findall(r"\d+", f.stem)[-1]) if re.findall(r"\d+", f.stem) else 0)
    return [f.read_bytes() for f in files]


def _rgb(fill, default):
    from PySide6.QtGui import QColor
    try:
        if fill.type == 1:           # MSO_FILL.SOLID
            return QColor("#" + str(fill.fore_color.rgb))
    except (AttributeError, TypeError, ValueError, KeyError):
        pass
    return default


def _font_color(run, default):
    from PySide6.QtGui import QColor
    try:
        if run.font.color and run.font.color.type is not None and run.font.color.rgb is not None:
            return QColor("#" + str(run.font.color.rgb))
    except (AttributeError, TypeError, ValueError):
        pass
    return default


def render_basic(pptx: Path, width: int = SLIDE_WIDTH) -> list[bytes]:
    """Tự dựng ảnh slide (không cần PowerPoint): nền, hình (chữ nhật / bo góc / tròn), ảnh, chữ, bảng."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Pt
    from PySide6.QtCore import QBuffer, QByteArray, QIODevice, QRectF, Qt
    from PySide6.QtGui import QColor, QFont, QFontDatabase, QFontMetricsF, QGuiApplication, QImage, QPainter, QPen

    global _QT_APP
    if QGuiApplication.instance() is None:   # cần để vẽ chữ khi chạy ngoài giao diện; dùng QApplication để
        from PySide6.QtWidgets import QApplication   # sau đó vẫn tạo được cửa sổ / widget trong cùng tiến trình
        _QT_APP = QApplication([])
    families = set(QFontDatabase.families())
    prs = Presentation(str(pptx))
    sw, sh = prs.slide_width or 12192000, prs.slide_height or 6858000
    scale = width / sw
    height = int(sh * scale)
    images = []

    def qfont(run, pt, bold, italic):
        f = QFont()
        name = run.font.name if run is not None else None
        if name and name in families:
            f.setFamily(name)
        f.setPixelSize(max(9, int(Pt(pt) * scale)))
        f.setBold(bool(bold))
        f.setItalic(bool(italic))
        return f

    def draw_text(p: QPainter, shape, rect: QRectF, is_title: bool):
        tf = shape.text_frame
        ml = (tf.margin_left if tf.margin_left is not None else 91440) * scale
        mr = (tf.margin_right if tf.margin_right is not None else 91440) * scale
        mt = (tf.margin_top if tf.margin_top is not None else 45720) * scale
        mb = (tf.margin_bottom if tf.margin_bottom is not None else 45720) * scale
        inner = rect.adjusted(ml, mt, -mr, -mb)
        blocks = []
        for para in tf.paragraphs:
            run = next((r for r in para.runs if r.text.strip()), para.runs[0] if para.runs else None)
            size = next((r.font.size for r in para.runs if r.font.size), None)
            pt = size.pt if size else (36 if is_title else 18)
            bold = is_title or any(r.font.bold for r in para.runs)
            italic = any(r.font.italic for r in para.runs)
            font = qfont(run, pt, bold, italic)
            color = _font_color(run, QColor("#1E2A1D")) if run is not None else QColor("#1E2A1D")
            align = {PP_ALIGN.CENTER: Qt.AlignHCenter, PP_ALIGN.RIGHT: Qt.AlignRight}.get(para.alignment, Qt.AlignLeft)
            text = para.text
            if para.level and not text.startswith("•"):
                text = "    " * para.level + text
            p.setFont(font)
            h = p.boundingRect(QRectF(0, 0, inner.width(), 10000), Qt.TextWordWrap | align, text or " ").height()
            if h <= 0:      # một số font đậm cỡ nhỏ trả về 0: ước lượng theo số dòng
                fm = QFontMetricsF(font)
                lines = max(1, int(fm.horizontalAdvance(text) // max(inner.width(), 1)) + 1)
                h = fm.lineSpacing() * lines
            blocks.append((text, font, color, align, h))
        total = sum(b[4] for b in blocks) + 4 * max(0, len(blocks) - 1)
        anchor = tf.vertical_anchor if tf.vertical_anchor is not None else (
            MSO_ANCHOR.MIDDLE if getattr(shape, "auto_shape_type", None) not in (None,) and not
            getattr(shape, "is_placeholder", False) and shape.shape_type == 1 else MSO_ANCHOR.TOP)
        y = inner.y()
        if anchor == MSO_ANCHOR.MIDDLE:
            y = inner.y() + (inner.height() - total) / 2
        elif anchor == MSO_ANCHOR.BOTTOM:
            y = inner.bottom() - total
        for text, font, color, align, h in blocks:
            p.setFont(font)
            p.setPen(color)
            p.drawText(QRectF(inner.x(), y, inner.width(), h), Qt.TextWordWrap | align, text)
            y += h + 4

    def draw_shape(p: QPainter, shape):
        if shape.left is None or shape.width is None:
            if shape.shape_type == 6:
                for s in shape.shapes:
                    draw_shape(p, s)
            return
        rect = QRectF(shape.left * scale, shape.top * scale, shape.width * scale, shape.height * scale)
        if shape.shape_type == 6:
            for s in shape.shapes:
                draw_shape(p, s)
            return
        if shape.shape_type == 13 and hasattr(shape, "image"):          # ảnh
            img = QImage.fromData(shape.image.blob)
            if not img.isNull():
                p.drawImage(rect, img)
            return
        try:
            fill = _rgb(shape.fill, None)
        except (AttributeError, TypeError):
            fill = None
        if fill is not None:
            p.setPen(Qt.NoPen)
            p.setBrush(fill)
            kind = getattr(shape, "auto_shape_type", None) if shape.shape_type == 1 else None
            if kind == MSO_SHAPE.OVAL:
                p.drawEllipse(rect)
            elif kind == MSO_SHAPE.ROUNDED_RECTANGLE:
                r = min(rect.width(), rect.height()) * 0.12
                p.drawRoundedRect(rect, r, r)
            else:
                p.drawRect(rect)
            p.setBrush(Qt.NoBrush)
        if getattr(shape, "has_table", False) and shape.has_table:
            rows = list(shape.table.rows)
            cols = len(rows[0].cells) if rows else 0
            if rows and cols:
                rh, cw = rect.height() / len(rows), rect.width() / cols
                f = QFont()
                f.setPixelSize(max(12, int(rh * 0.35)))
                p.setFont(f)
                for r, row in enumerate(rows):
                    for c, cell in enumerate(row.cells):
                        cr = QRectF(rect.x() + c * cw, rect.y() + r * rh, cw, rh)
                        if r == 0:
                            p.fillRect(cr, QColor("#ECEFE3"))
                        p.setPen(QPen(QColor("#9CA38A"), 1))
                        p.drawRect(cr)
                        p.setPen(QColor("#1E2A1D"))
                        p.drawText(cr.adjusted(8, 4, -8, -4), Qt.AlignVCenter | Qt.TextWordWrap, cell.text)
            return
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip():
            is_title = bool(shape.is_placeholder and shape.placeholder_format.type in (1, 3))
            draw_text(p, shape, rect, is_title)

    for slide in prs.slides:
        img = QImage(width, height, QImage.Format_RGB32)
        img.fill(QColor("white"))
        p = QPainter(img)
        p.setRenderHint(QPainter.Antialiasing)
        p.setRenderHint(QPainter.TextAntialiasing)
        try:
            color = _rgb(slide.background.fill, None)
            if color is not None:
                p.fillRect(img.rect(), color)
        except (AttributeError, TypeError):
            pass
        for shape in slide.shapes:
            try:
                draw_shape(p, shape)
            except Exception:  # hình lạ: bỏ qua, không làm hỏng cả slide
                continue
        p.end()
        buf = QByteArray()
        dev = QBuffer(buf)
        dev.open(QIODevice.WriteOnly)
        img.save(dev, "PNG")
        images.append(bytes(buf.data()))
    return images


# ====================================================================== nhập bài giảng


def import_pptx(pptx: Path, ten: str, mon: str, chuong: int | None = None, ten_en: str = "",
                mo_ta: str = "", dest_base: Path | None = None, use_powerpoint: bool = True) -> Lesson:
    """Nhập một file .pptx thành bài giảng (ảnh slide đóng gói + chữ để tra cứu)."""
    pptx = Path(pptx)
    slides = extract_text(pptx)
    images: list[bytes] = []
    if use_powerpoint:
        with tempfile.TemporaryDirectory() as tmp:
            images = export_with_powerpoint(pptx.resolve(), Path(tmp))
    if len(images) != len(slides):
        images = render_basic(pptx)
    folder, lesson_id = _new_folder(ten, dest_base, pptx)
    pack_slides(folder, lesson_id, images)
    return _write_meta(folder, {"id": lesson_id, "loai": "slides", "ten": ten, "ten_en": ten_en,
                                "mon": mon.upper(), "chuong": chuong, "mo_ta": mo_ta, "nguon": pptx.name,
                                "slides": slides})


def _new_folder(ten: str, dest_base, source: Path) -> tuple[Path, str]:
    base = Path(dest_base) if dest_base else editable_dir()
    slug = _slug(ten)
    folder, n = base / slug, 1
    while folder.exists():
        n += 1
        folder = base / f"{slug}_{n}"
    folder.mkdir(parents=True)
    h = hashlib.sha1()
    if source.is_file():
        with open(source, "rb") as f:
            for chunk in iter(lambda: f.read(1 << 20), b""):
                h.update(chunk)
    else:
        h.update(str(source).encode("utf-8"))
    return folder, f"{slug}-{h.hexdigest()[:10]}"


def _write_meta(folder: Path, data: dict) -> Lesson:
    (folder / META).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    lesson = load_lesson(folder)
    if lesson is None:
        shutil.rmtree(folder, ignore_errors=True)
        raise ValueError("bài giảng không hợp lệ")
    return lesson


def import_video(video: Path, ten: str, mon: str, chuong: int | None = None, ten_en: str = "",
                 mo_ta: str = "", dest_base: Path | None = None) -> Lesson:
    """Nhập một video ngắn: mã hóa và cất trong thư mục bài giảng (không chép file gốc)."""
    video = Path(video)
    if video.suffix.lower() not in VIDEO_EXT:
        raise ValueError(f"không phải file video ({', '.join(VIDEO_EXT)})")
    folder, lesson_id = _new_folder(ten, dest_base, video)
    (folder / VIDEO).write_bytes(_xor(video.read_bytes(), _key(lesson_id)))
    return _write_meta(folder, {"id": lesson_id, "loai": "video", "ten": ten, "ten_en": ten_en,
                                "mon": mon.upper(), "chuong": chuong, "mo_ta": mo_ta, "nguon": video.name,
                                "dinh_dang": video.suffix.lower(),
                                "slides": [{"text": f"{ten}\n{ten_en}\n{mo_ta}".strip(), "notes": ""}]})


# ---------------------------------------------------------------- SCORM


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def read_manifest(xml: bytes) -> dict:
    """Đọc imsmanifest.xml → {"launch": trang mở đầu, "phien_ban": "1.2"/"2004", "tieu_de": [...]}."""
    from xml.etree import ElementTree as ET
    root = ET.fromstring(xml)
    text = xml.decode("utf-8", "ignore")
    version = "2004" if ("2004" in text and "adlcp_v1p3" in text) or "CAM 1.3" in text else "1.2"
    for el in root.iter():
        if _local(el.tag) == "schemaversion" and el.text:
            version = "2004" if ("2004" in el.text or "1.3" in el.text) else "1.2"
    resources = {}
    for el in root.iter():
        if _local(el.tag) == "resource" and el.get("href"):
            base = next((v for k, v in el.attrib.items() if _local(k) == "base"), "")
            resources[el.get("identifier")] = (base + el.get("href")).split("?")[0]
    titles, launch = [], ""
    for el in root.iter():
        if _local(el.tag) == "title" and el.text and el.text.strip():
            titles.append(el.text.strip())
        if _local(el.tag) == "item" and not launch and el.get("identifierref") in resources:
            launch = resources[el.get("identifierref")]
    if not launch and resources:
        launch = next(iter(resources.values()))
    if not launch:
        raise ValueError("imsmanifest.xml không có trang mở đầu (resource href)")
    return {"launch": launch, "phien_ban": version, "tieu_de": titles}


def import_scorm(package: Path, ten: str, mon: str, chuong: int | None = None, ten_en: str = "",
                 mo_ta: str = "", dest_base: Path | None = None) -> Lesson:
    """Nhập gói SCORM (.zip hoặc thư mục có imsmanifest.xml)."""
    package = Path(package)
    if package.is_dir():
        man = package / "imsmanifest.xml"
        if not man.is_file():
            raise ValueError("thư mục không có imsmanifest.xml")
        info = read_manifest(man.read_bytes())
        folder, lesson_id = _new_folder(ten, dest_base, package)
        shutil.copytree(package, folder / SCORM_DIR)
    else:
        with zipfile.ZipFile(package) as z:
            names = z.namelist()
            man = next((n for n in names if n.rsplit("/", 1)[-1].lower() == "imsmanifest.xml"), None)
            if man is None:
                raise ValueError("file .zip không phải gói SCORM (thiếu imsmanifest.xml)")
            info = read_manifest(z.read(man))
            prefix = man[:-len("imsmanifest.xml")]
            folder, lesson_id = _new_folder(ten, dest_base, package)
            dest = (folder / SCORM_DIR).resolve()
            for n in names:
                if not n.startswith(prefix) or n.endswith("/"):
                    continue
                target = (dest / n[len(prefix):]).resolve()
                if dest not in target.parents:       # chặn đường dẫn "../" trong zip
                    continue
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(z.read(n))
    text = "\n".join([ten, ten_en, mo_ta] + info["tieu_de"]).strip()
    return _write_meta(folder, {"id": lesson_id, "loai": "scorm", "ten": ten, "ten_en": ten_en,
                                "mon": mon.upper(), "chuong": chuong, "mo_ta": mo_ta, "nguon": package.name,
                                "launch": info["launch"], "phien_ban": info["phien_ban"],
                                "slides": [{"text": text, "notes": ""}]})


def import_file(path: Path, ten: str, mon: str, chuong: int | None = None, ten_en: str = "",
                mo_ta: str = "", dest_base: Path | None = None) -> Lesson:
    """Nhập bài giảng theo đuôi file: .pptx → slide, video → video, .zip / thư mục → SCORM."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pptx":
        return import_pptx(path, ten, mon, chuong, ten_en, mo_ta, dest_base)
    if ext in VIDEO_EXT:
        return import_video(path, ten, mon, chuong, ten_en, mo_ta, dest_base)
    if ext == ".zip" or path.is_dir():
        return import_scorm(path, ten, mon, chuong, ten_en, mo_ta, dest_base)
    raise ValueError("chỉ nhận .pptx, video (.mp4…) hoặc gói SCORM (.zip)")


# ====================================================================== tiến độ học


def _progress_file() -> Path:
    return DATA_DIR / "tien_do_hoc.json"


def load_progress(user: str | None) -> dict[str, int]:
    """{id bài giảng: slide xa nhất đã xem} của một tài khoản."""
    try:
        data = json.loads(_progress_file().read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return data.get(user or "", {})


def save_progress(user: str | None, lesson_id: str, slide: int) -> bool:
    """Ghi slide xa nhất đã xem; True nếu tiến độ tăng."""
    try:
        data = json.loads(_progress_file().read_text(encoding="utf-8"))
    except (OSError, ValueError):
        data = {}
    mine = data.setdefault(user or "", {})
    if slide > mine.get(lesson_id, 0):
        mine[lesson_id] = slide
        _progress_file().parent.mkdir(parents=True, exist_ok=True)
        _progress_file().write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        return True
    return False


# ====================================================================== dữ liệu SCORM (cmi.*) của học viên


def _scorm_file() -> Path:
    return DATA_DIR / "scorm_hoc_vien.json"


def load_scorm(user: str | None, lesson_id: str) -> dict[str, str]:
    try:
        data = json.loads(_scorm_file().read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return dict(data.get(user or "", {}).get(lesson_id, {}))


def save_scorm(user: str | None, lesson_id: str, cmi: dict[str, str]) -> None:
    """Lưu dữ liệu cmi.* mà bài SCORM gửi về; hoàn thành / đạt → tính là đã học."""
    try:
        data = json.loads(_scorm_file().read_text(encoding="utf-8"))
    except (OSError, ValueError):
        data = {}
    data.setdefault(user or "", {})[lesson_id] = {k: str(v) for k, v in cmi.items()}
    _scorm_file().parent.mkdir(parents=True, exist_ok=True)
    _scorm_file().write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    if scorm_done(cmi):
        save_progress(user, lesson_id, 1)


def scorm_done(cmi: dict) -> bool:
    status = {str(cmi.get(k, "")).lower() for k in ("cmi.core.lesson_status", "cmi.completion_status",
                                                     "cmi.success_status")}
    return bool(status & {"completed", "passed"})


def scorm_score(cmi: dict) -> str:
    """Điểm bài SCORM dạng chữ (vd "80/100"), rỗng nếu chưa có."""
    raw = cmi.get("cmi.core.score.raw") or cmi.get("cmi.score.raw")
    if raw in (None, ""):
        scaled = cmi.get("cmi.score.scaled")
        return f"{round(float(scaled) * 100)}%" if scaled not in (None, "") else ""
    top = cmi.get("cmi.core.score.max") or cmi.get("cmi.score.max") or "100"
    return f"{raw}/{top}"
