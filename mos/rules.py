"""Thư viện "luật chấm" dùng cho đề tự soạn (de.json).

Mỗi luật là một hàm `(path, **tham_so) -> bool` đọc file người học đã lưu.
Trong de.json, một nhiệm vụ khai báo:

    "cham": {"luat": "excel_cong_thuc", "o": "E2:E11", "chua": ["C{hang}", "*", "D{hang}"]}

hoặc một danh sách luật (phải đúng TẤT CẢ):

    "cham": [{"luat": "excel_ten_sheet", "ten": "DoanhSo"}, {"luat": "excel_co_dinh", "o": "A2"}]

Xem danh sách đầy đủ trong HUONG_DAN_SOAN_DE.md hoặc chạy:  py kiem_tra_de.py --luat
"""
from __future__ import annotations

import inspect
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from pptx import Presentation

from .ooxml import any_part_contains, norm, part, parts

RULES: dict = {}


def rule(fn):
    """Đăng ký một luật chấm; tên hàm = tên luật dùng trong de.json."""
    RULES[fn.__name__] = fn
    return fn


def _as_list(x):
    return x if isinstance(x, list) else [x]


def _compact(text) -> str:
    return re.sub(r"\s+", "", str(text or "")).upper()


# ====================================================================== EXCEL


def _ws(path, sheet=None):
    wb = load_workbook(path)
    if sheet:
        match = [n for n in wb.sheetnames if n.casefold() == str(sheet).casefold()]
        if not match:
            raise KeyError(f"không có trang tính '{sheet}'")
        return wb, wb[match[0]]
    return wb, wb.active


def _cells(ws, o):
    """'E2' hoặc 'E2:E11' → danh sách ô."""
    min_col, min_row, max_col, max_row = range_boundaries(o)
    return [ws.cell(row=r, column=c) for r in range(min_row, max_row + 1)
            for c in range(min_col, max_col + 1)]


def _formula(cell) -> str:
    v = cell.value
    v = getattr(v, "text", v)
    return _compact(v) if isinstance(v, str) and v.startswith("=") else ""


@rule
def excel_ten_sheet(path, ten):
    """Có trang tính tên `ten`."""
    return any(n.casefold() == ten.casefold() for n in load_workbook(path).sheetnames)


@rule
def excel_khong_co_sheet(path, ten):
    """KHÔNG còn trang tính tên `ten` (vd đã đổi tên / xóa)."""
    return not excel_ten_sheet(path, ten)


@rule
def excel_cong_thuc(path, o, chua, sheet=None):
    """Mọi ô trong vùng `o` là công thức chứa tất cả chuỗi trong `chua`.
    Dùng {hang} để thay bằng số hàng của từng ô, vd "C{hang}*D{hang}"."""
    _, ws = _ws(path, sheet)
    for cell in _cells(ws, o):
        f = _formula(cell)
        if not f or not all(_compact(s.replace("{hang}", str(cell.row))) in f for s in _as_list(chua)):
            return False
    return True


@rule
def excel_gia_tri(path, o, bang, sheet=None):
    """Ô `o` có giá trị bằng `bang` (so sánh không phân biệt hoa/thường)."""
    _, ws = _ws(path, sheet)
    v = ws[o].value
    if isinstance(bang, (int, float)) and isinstance(v, (int, float)):
        return abs(v - bang) < 1e-9
    return norm(str(v if v is not None else "")) == norm(str(bang))


@rule
def excel_dinh_dang_so(path, o, chua=None, sheet=None):
    """Vùng `o` có định dạng số khác General; nếu có `chua` thì mã định dạng phải chứa chuỗi đó (vd "%", "0.00")."""
    _, ws = _ws(path, sheet)
    for cell in _cells(ws, o):
        fmt = cell.number_format
        if fmt in ("General", "@") or (chua and chua not in fmt):
            return False
    return True


@rule
def excel_font(path, o, dam=None, nghieng=None, co=None, sheet=None):
    """Font của vùng `o`: dam (in đậm), nghieng, co (cỡ chữ)."""
    _, ws = _ws(path, sheet)
    for cell in _cells(ws, o):
        f = cell.font
        if dam is not None and bool(f.b) != dam:
            return False
        if nghieng is not None and bool(f.i) != nghieng:
            return False
        if co is not None and (f.sz or 11) != co:
            return False
    return True


@rule
def excel_co_bang(path, vung, sheet=None, ten=None):
    """Có Table (Format as Table) bắt đầu ở góc trên-trái và phủ hết `vung`."""
    _, ws = _ws(path, sheet)
    want = range_boundaries(vung)
    for t in ws.tables.values():
        got = range_boundaries(t.ref)
        if got[:3] == want[:3] and got[3] >= want[3] and (not ten or t.displayName.casefold() == ten.casefold()):
            return True
    return False


@rule
def excel_co_dinh(path, o="A2", sheet=None):
    """Freeze Panes tại ô `o` (A2 = cố định hàng đầu, B1 = cột đầu)."""
    return _ws(path, sheet)[1].freeze_panes == o


@rule
def excel_dinh_dang_dieu_kien(path, vung, sheet=None):
    """Có Conditional Formatting phủ vùng `vung`."""
    _, ws = _ws(path, sheet)
    min_col, min_row, max_col, max_row = range_boundaries(vung)
    corners = [ws.cell(row=r, column=c).coordinate for r in (min_row, max_row) for c in (min_col, max_col)]
    return any(cf.rules and all(x in cf.sqref for x in corners) for cf in ws.conditional_formatting)


@rule
def excel_ten_vung(path, ten, vung):
    """Có Named Range `ten` trỏ tới `vung` (vd "C2:C11")."""
    wb = load_workbook(path)
    want = ":".join("$%s$%s" % re.fullmatch(r"([A-Z]+)(\d+)", c.upper()).groups() for c in vung.split(":"))
    names = list(wb.defined_names.items())
    for ws in wb.worksheets:
        names += list(ws.defined_names.items())
    return any(key.casefold() == ten.casefold() and dn.attr_text.replace("'", "").upper().endswith(want)
               for key, dn in names)


@rule
def excel_bieu_do(path, loai="bat_ky"):
    """Có biểu đồ. loai: cot, thanh (ngang), duong, tron, bat_ky."""
    pattern = {
        "cot": r'<(\w+:)?barDir val="col"',
        "thanh": r'<(\w+:)?barDir val="bar"',
        "duong": r"<(\w+:)?lineChart\b",
        "tron": r"<(\w+:)?(pie|doughnut|pie3D)Chart\b",
        "bat_ky": r"<(\w+:)?plotArea\b",
    }[loai]
    return any_part_contains(path, r"xl/charts/chart\d+\.xml", pattern)


@rule
def excel_huong_trang(path, huong="ngang", sheet=None):
    """Hướng trang in: ngang hoặc doc."""
    want = {"ngang": "landscape", "doc": "portrait"}[huong]
    return _ws(path, sheet)[1].page_setup.orientation == want


@rule
def excel_sap_xep(path, cot, tu, den, chieu="tang", sheet=None):
    """Cột `cot` từ hàng `tu` đến `den` được sắp xếp tang / giam."""
    _, ws = _ws(path, sheet)
    vals = [ws[f"{cot}{r}"].value for r in range(tu, den + 1)]
    try:
        return vals == sorted(vals, reverse=(chieu == "giam"))
    except TypeError:
        return False


# ====================================================================== WORD


def _para(doc, doan):
    target = norm(doan)
    for p in doc.paragraphs:
        if norm(p.text).startswith(target):
            return p
    raise KeyError(f"không tìm thấy đoạn bắt đầu bằng '{doan}'")


def _doc_text(doc) -> str:
    chunks = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            chunks.extend(c.text for c in row.cells)
    return norm("\n".join(chunks))


FIELD = r"(instrText[^>]*>\s*{0}\b|w:instr=\"\s*{0}\b)"


@rule
def word_kieu_doan(path, doan, kieu):
    """Các đoạn (bắt đầu bằng `doan`, có thể là danh sách) dùng style `kieu` (vd "Heading 1", "Title")."""
    doc = Document(path)
    return all(_para(doc, d).style.name.casefold() == kieu.casefold() for d in _as_list(doan))


@rule
def word_can_le(path, doan, can="giua"):
    """Căn lề đoạn: trai, giua, phai, deu."""
    want = {"trai": WD_ALIGN_PARAGRAPH.LEFT, "giua": WD_ALIGN_PARAGRAPH.CENTER,
            "phai": WD_ALIGN_PARAGRAPH.RIGHT, "deu": WD_ALIGN_PARAGRAPH.JUSTIFY}[can]
    doc = Document(path)
    for d in _as_list(doan):
        p = _para(doc, d)
        align = p.alignment if p.alignment is not None else p.style.paragraph_format.alignment
        if align != want:
            return False
    return True


@rule
def word_co_chu(path, chu, so_lan=None):
    """Tài liệu có chứa `chu` (nếu có so_lan thì phải xuất hiện đúng số lần)."""
    text, c = _doc_text(Document(path)), norm(chu)
    return text.count(c) == so_lan if so_lan is not None else c in text


@rule
def word_khong_co_chu(path, chu):
    """Tài liệu KHÔNG còn chứa `chu` (vd sau Replace All)."""
    return norm(chu) not in _doc_text(Document(path))


@rule
def word_dau_dong(path, doan):
    """Các đoạn trong `doan` là danh sách bullet/đánh số."""
    doc = Document(path)
    for d in _as_list(doan):
        p = _para(doc, d)
        has_num = p._p.pPr is not None and p._p.pPr.numPr is not None
        if not (has_num or p.style.name.startswith("List")):
            return False
    return True


@rule
def word_gian_dong(path, doan, gia_tri):
    """Giãn dòng của đoạn = gia_tri (vd 1.5, 2)."""
    p = _para(Document(path), doan)
    ls = p.paragraph_format.line_spacing or p.style.paragraph_format.line_spacing
    return isinstance(ls, float) and abs(ls - gia_tri) < 0.01


@rule
def word_bang(path, so_cot, so_hang=None, o_dau=None):
    """Có bảng `so_cot` cột (và ít nhất `so_hang` hàng; ô đầu = `o_dau` nếu có)."""
    for t in Document(path).tables:
        if len(t.columns) != so_cot or (so_hang and len(t.rows) < so_hang):
            continue
        if o_dau and norm(t.cell(0, 0).text) != norm(o_dau):
            continue
        return True
    return False


@rule
def word_muc_luc(path):
    """Có mục lục tự động (Table of Contents)."""
    return any_part_contains(path, r"word/document\.xml", FIELD.format("TOC"))


@rule
def word_so_trang(path, vi_tri="bat_ky"):
    """Có số trang (trường PAGE) ở chan_trang, dau_trang hoặc bat_ky."""
    pat = {"chan_trang": r"word/footer\d*\.xml", "dau_trang": r"word/header\d*\.xml",
           "bat_ky": r"word/(header|footer)\d*\.xml"}[vi_tri]
    return any_part_contains(path, pat, FIELD.format("PAGE"))


@rule
def word_dau_trang_chan_trang(path, chu, vi_tri="bat_ky"):
    """Đầu/chân trang có chứa `chu`."""
    pat = {"chan_trang": r"word/footer\d*\.xml", "dau_trang": r"word/header\d*\.xml",
           "bat_ky": r"word/(header|footer)\d*\.xml"}[vi_tri]
    return any(norm(chu) in norm(re.sub(r"<[^>]+>", "", x)) for x in parts(path, pat).values())


@rule
def word_watermark(path, chu):
    """Có watermark chữ `chu`."""
    return any_part_contains(path, r"word/header\d*\.xml", re.escape(chu), re.IGNORECASE)


@rule
def word_huong_trang(path, huong="ngang"):
    """Hướng trang của section đầu: ngang hoặc doc."""
    s = Document(path).sections[0]
    if huong == "ngang":
        return s.orientation == WD_ORIENT.LANDSCAPE and s.page_width > s.page_height
    return s.page_width < s.page_height


@rule
def word_so_cot(path, so):
    """Có đoạn văn bản được chia `so` cột (Layout > Columns)."""
    return any_part_contains(path, r"word/document\.xml", rf'<w:cols [^>]*w:num="{so}"')


@rule
def word_footnote(path, chu):
    """Có footnote chứa `chu`."""
    return norm(chu) in norm(re.sub(r"<[^>]+>", "", part(path, "word/footnotes.xml")))


@rule
def word_thuoc_tinh(path, truong, gia_tri):
    """Thuộc tính tài liệu: truong = title, author, subject, keywords, comments, category."""
    return norm(str(getattr(Document(path).core_properties, truong) or "")) == norm(gia_tri)


@rule
def word_theo_doi_thay_doi(path):
    """Đã bật Track Changes."""
    m = re.search(r"<w:trackRevisions(?:\s+w:val=\"(\w+)\")?\s*/>", part(path, "word/settings.xml"))
    return bool(m) and (m.group(1) or "true") not in ("0", "false", "off")


@rule
def word_hinh_anh(path, alt_text=None):
    """Có hình ảnh (nếu có alt_text thì mô tả thay thế phải chứa chuỗi này)."""
    xml = part(path, "word/document.xml")
    if "<pic:pic" not in xml and "<w:drawing" not in xml:
        return False
    return alt_text is None or norm(alt_text) in norm(" ".join(re.findall(r'descr="([^"]*)"', xml)))


# ====================================================================== POWERPOINT


def _title(slide) -> str:
    t = slide.shapes.title
    return t.text_frame.text if t is not None and t.has_text_frame else ""


def _slide(prs, tieu_de=None, slide=None):
    if slide is not None:
        return prs.slides[slide - 1]
    for s in prs.slides:
        if norm(_title(s)) == norm(tieu_de):
            return s
    raise KeyError(f"không có slide tiêu đề '{tieu_de}'")


def _slides(prs, tieu_de=None, slide=None):
    if tieu_de is None and slide is None:
        return list(prs.slides)
    return [_slide(prs, tieu_de, slide)]


@rule
def ppt_co_slide(path, tieu_de, bo_cuc=None, vi_tri=None):
    """Có slide tiêu đề `tieu_de` (bố cục `bo_cuc`, ở vị trí `vi_tri`: số thứ tự hoặc "cuoi")."""
    prs = Presentation(path)
    slides = list(prs.slides)
    for i, s in enumerate(slides, start=1):
        if norm(_title(s)) != norm(tieu_de):
            continue
        if bo_cuc and s.slide_layout.name.casefold() != bo_cuc.casefold():
            continue
        if vi_tri == "cuoi" and i != len(slides):
            continue
        if isinstance(vi_tri, int) and i != vi_tri:
            continue
        return True
    return False


@rule
def ppt_khong_co_slide(path, tieu_de):
    """Đã xóa slide tiêu đề `tieu_de`."""
    return all(norm(_title(s)) != norm(tieu_de) for s in Presentation(path).slides)


@rule
def ppt_so_slide(path, so):
    """Bài có đúng `so` slide."""
    return len(Presentation(path).slides) == so


@rule
def ppt_ghi_chu(path, chu, tieu_de=None, slide=None):
    """Slide có Notes chứa `chu`."""
    s = _slide(Presentation(path), tieu_de, slide)
    return s.has_notes_slide and norm(chu) in norm(s.notes_slide.notes_text_frame.text)


@rule
def ppt_chuyen_trang(path, tieu_de=None, slide=None):
    """Có Transition (không ghi tieu_de/slide = tất cả slide)."""
    return all("<p:transition" in s._element.xml for s in _slides(Presentation(path), tieu_de, slide))


@rule
def ppt_hieu_ung(path, tieu_de=None, slide=None, nhom="entr"):
    """Slide có Animation. nhom: entr (xuất hiện), exit (biến mất), emph (nhấn mạnh), bat_ky."""
    xml = _slide(Presentation(path), tieu_de, slide)._element.xml
    return "<p:timing" in xml and (nhom == "bat_ky" or f'presetClass="{nhom}"' in xml)


@rule
def ppt_an_slide(path, tieu_de=None, slide=None):
    """Slide bị ẩn (Hide Slide)."""
    return _slide(Presentation(path), tieu_de, slide)._element.get("show") in ("0", "false")


@rule
def ppt_kich_thuoc(path, ti_le="16:9"):
    """Kích thước slide: "16:9" hoặc "4:3"."""
    prs = Presentation(path)
    a, b = (int(x) for x in ti_le.split(":"))
    return abs(prs.slide_width / prs.slide_height - a / b) < 0.01


@rule
def ppt_bang(path, so_cot, so_hang, tieu_de=None, slide=None):
    """Slide có bảng `so_cot` cột × `so_hang` hàng."""
    s = _slide(Presentation(path), tieu_de, slide)
    return any(sh.has_table and len(sh.table.columns) == so_cot and len(sh.table.rows) == so_hang
               for sh in s.shapes)


@rule
def ppt_smartart(path, tieu_de=None, slide=None):
    """Slide có SmartArt."""
    return "drawingml/2006/diagram" in _slide(Presentation(path), tieu_de, slide)._element.xml


@rule
def ppt_bieu_do(path, tieu_de=None, slide=None):
    """Slide có biểu đồ (Chart)."""
    return any(getattr(sh, "has_chart", False) for sh in _slide(Presentation(path), tieu_de, slide).shapes)


@rule
def ppt_hinh_anh(path, tieu_de=None, slide=None):
    """Slide có hình ảnh."""
    return "<p:pic" in _slide(Presentation(path), tieu_de, slide)._element.xml


@rule
def ppt_section(path, ten):
    """Có Section tên `ten`."""
    return re.search(rf'<p14:section [^>]*name="{re.escape(ten)}"', part(path, "ppt/presentation.xml")) is not None


@rule
def ppt_so_trang(path, tru_slide_dau=True):
    """Hiện Slide number trên mọi slide (mặc định bỏ qua slide tiêu đề)."""
    slides = list(Presentation(path).slides)[1 if tru_slide_dau else 0:]
    return bool(slides) and all(re.search(r'<p:ph [^>]*type="sldNum"', s._element.xml) for s in slides)


@rule
def ppt_co_chu(path, chu, tieu_de=None, slide=None):
    """Slide (hoặc cả bài) có chứa `chu`."""
    for s in _slides(Presentation(path), tieu_de, slide):
        text = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        if norm(chu) in norm(text):
            return True
    return False


# ====================================================================== BỘ ĐỀ NHẬP


@rule
def word_mau_de(path, dang, de_bai, goc=None):
    """Chấm tự động theo dạng câu của bộ đề nhập (mos/word_auto.py): `dang` là mã dạng
    (vd pic_size), `de_bai` là đề tiếng Anh, `goc` là thông tin file gốc. Dạng chưa hỗ trợ
    thì không tính điểm (học viên tự kiểm tra)."""
    from . import word_auto
    return word_auto.grade(path, dang, de_bai, goc)


@rule
def tu_kiem_tra(path, ghi_chu=None):
    """Không chấm tự động: học viên tự đối chiếu với gợi ý, nhiệm vụ không tính vào điểm."""
    return None


# ====================================================================== NÂNG CAO


@rule
def xml_chua(path, part_regex, chua, bo_qua_hoa_thuong=False):
    """Luật nâng cao: có part XML (khớp regex `part_regex`) chứa regex `chua`."""
    return any_part_contains(path, part_regex, chua, re.IGNORECASE if bo_qua_hoa_thuong else 0)


# ====================================================================== thông tin cho form soạn đề

# tên tham số -> (nhãn, kiểu, lựa chọn). Kiểu: text, number, auto, list, bool, choice
PARAM_UI = {
    "ten": ("Tên", "text"),
    "o": ("Ô / vùng (vd E2 hoặc E2:E11)", "text"),
    "vung": ("Vùng (vd A1:E11)", "text"),
    "sheet": ("Trang tính (để trống = trang đang chọn)", "text"),
    "chua": ("Phải chứa", "text"),
    "bang": ("Giá trị đúng", "auto"),
    "dam": ("In đậm", "bool"),
    "nghieng": ("In nghiêng", "bool"),
    "co": ("Cỡ chữ", "number"),
    "loai": ("Loại biểu đồ", "choice", ["bat_ky", "cot", "thanh", "duong", "tron"]),
    "huong": ("Hướng trang", "choice", ["ngang", "doc"]),
    "cot": ("Cột (vd C)", "text"),
    "tu": ("Từ hàng", "number"),
    "den": ("Đến hàng", "number"),
    "chieu": ("Chiều sắp xếp", "choice", ["tang", "giam"]),
    "doan": ("Đoạn văn bắt đầu bằng (nhiều đoạn: cách nhau bởi ;)", "list"),
    "kieu": ("Style (vd Heading 1, Title)", "text"),
    "can": ("Căn lề", "choice", ["trai", "giua", "phai", "deu"]),
    "chu": ("Nội dung chữ", "text"),
    "so_lan": ("Số lần xuất hiện", "number"),
    "gia_tri": ("Giá trị", "auto"),
    "so_cot": ("Số cột", "number"),
    "so_hang": ("Số hàng", "number"),
    "o_dau": ("Chữ trong ô đầu tiên", "text"),
    "vi_tri": ("Vị trí", "choice", ["bat_ky", "chan_trang", "dau_trang"]),
    "truong": ("Thuộc tính", "choice", ["title", "author", "subject", "keywords", "comments", "category"]),
    "so": ("Số lượng", "number"),
    "alt_text": ("Alt text phải chứa", "text"),
    "tieu_de": ("Tiêu đề slide", "text"),
    "slide": ("Hoặc số thứ tự slide", "number"),
    "bo_cuc": ("Bố cục (vd Title and Content)", "text"),
    "nhom": ("Nhóm hiệu ứng", "choice", ["entr", "exit", "emph", "bat_ky"]),
    "ti_le": ("Tỉ lệ", "choice", ["16:9", "4:3"]),
    "tru_slide_dau": ("Bỏ qua slide tiêu đề", "bool"),
    "dang": ("Mã dạng câu (vd pic_size)", "text"),
    "de_bai": ("Đề bài gốc (tiếng Anh)", "text"),
    "goc": ("Thông tin file gốc (tự sinh khi nhập đề)", "auto"),
    "ghi_chu": ("Ghi chú", "text"),
    "part_regex": ("Part XML (regex, vd word/document\\.xml)", "text"),
    "bo_qua_hoa_thuong": ("Không phân biệt hoa/thường", "bool"),
}
PARAM_UI_RULE = {
    ("excel_cong_thuc", "chua"): ("Công thức phải chứa (cách nhau bởi ; — {hang} = số hàng)", "list"),
    ("excel_dinh_dang_so", "chua"): ("Mã định dạng phải chứa (vd %, 0.00)", "text"),
    ("xml_chua", "chua"): ("Regex cần tìm", "text"),
    ("ppt_co_slide", "vi_tri"): ("Vị trí (số thứ tự, hoặc cuoi)", "auto"),
    ("word_gian_dong", "doan"): ("Đoạn văn bắt đầu bằng", "text"),
    ("word_gian_dong", "gia_tri"): ("Giãn dòng (vd 1.5)", "number"),
    ("word_thuoc_tinh", "gia_tri"): ("Giá trị", "text"),
}
BOOL_TEXT = {"Có": True, "Không": False}


RULE_TITLES = {
    "excel_ten_sheet": "Có trang tính với tên cho trước",
    "excel_khong_co_sheet": "Không còn trang tính có tên cho trước (đã đổi tên / xóa)",
    "excel_cong_thuc": "Ô / vùng chứa công thức đúng",
    "excel_gia_tri": "Ô có giá trị đúng",
    "excel_dinh_dang_so": "Định dạng số (Number Format) cho vùng",
    "excel_font": "Font chữ: in đậm / in nghiêng / cỡ chữ",
    "excel_co_bang": "Đã định dạng thành bảng (Format as Table)",
    "excel_co_dinh": "Cố định hàng / cột (Freeze Panes)",
    "excel_dinh_dang_dieu_kien": "Định dạng có điều kiện (Conditional Formatting)",
    "excel_ten_vung": "Đặt tên vùng (Named Range)",
    "excel_bieu_do": "Có biểu đồ (Chart)",
    "excel_huong_trang": "Hướng trang in (ngang / dọc)",
    "excel_sap_xep": "Dữ liệu được sắp xếp (Sort)",
    "word_kieu_doan": "Đoạn văn dùng Style (Heading 1, Title…)",
    "word_can_le": "Căn lề đoạn văn",
    "word_co_chu": "Tài liệu có chứa đoạn chữ",
    "word_khong_co_chu": "Tài liệu không còn đoạn chữ (vd sau Replace All)",
    "word_dau_dong": "Đoạn văn là danh sách (Bullets / Numbering)",
    "word_gian_dong": "Giãn dòng (Line Spacing)",
    "word_bang": "Có bảng (Table) với số cột / hàng",
    "word_muc_luc": "Có mục lục tự động (Table of Contents)",
    "word_so_trang": "Có số trang (Page Number)",
    "word_dau_trang_chan_trang": "Đầu trang / chân trang có chứa chữ",
    "word_watermark": "Có watermark chữ",
    "word_huong_trang": "Hướng trang (ngang / dọc)",
    "word_so_cot": "Chia cột văn bản (Columns)",
    "word_footnote": "Có chú thích cuối trang (Footnote)",
    "word_thuoc_tinh": "Thuộc tính tài liệu (Title, Author…)",
    "word_theo_doi_thay_doi": "Đã bật theo dõi thay đổi (Track Changes)",
    "word_hinh_anh": "Có hình ảnh (và Alt Text)",
    "ppt_co_slide": "Có slide với tiêu đề cho trước",
    "ppt_khong_co_slide": "Đã xóa slide có tiêu đề cho trước",
    "ppt_so_slide": "Bài có đúng số slide",
    "ppt_ghi_chu": "Slide có ghi chú (Notes)",
    "ppt_chuyen_trang": "Có hiệu ứng chuyển trang (Transition)",
    "ppt_hieu_ung": "Có hiệu ứng động (Animation)",
    "ppt_an_slide": "Slide bị ẩn (Hide Slide)",
    "ppt_kich_thuoc": "Kích thước slide (16:9 / 4:3)",
    "ppt_bang": "Slide có bảng với số cột / hàng",
    "ppt_smartart": "Slide có SmartArt",
    "ppt_bieu_do": "Slide có biểu đồ (Chart)",
    "ppt_hinh_anh": "Slide có hình ảnh",
    "ppt_section": "Có Section với tên cho trước",
    "ppt_so_trang": "Hiện số slide (Slide Number)",
    "ppt_co_chu": "Slide có chứa đoạn chữ",
    "word_mau_de": "Bộ đề nhập: chấm tự động theo dạng câu",
    "tu_kiem_tra": "Không chấm tự động (học viên tự kiểm tra)",
    "xml_chua": "Nâng cao: file chứa đoạn XML (regex)",
}


def rule_title(name: str) -> str:
    """Tên ngắn tiếng Việt của luật (hiện trong form soạn đề)."""
    from .i18n import tr
    return tr(RULE_TITLES.get(name, name))


def rules_for(mon: str) -> list[str]:
    prefix = {"EXCEL": "excel_", "WORD": "word_", "POWERPOINT": "ppt_"}.get(mon.upper(), "")
    return [n for n in RULES if n.startswith(prefix)] + [n for n in RULES if n.startswith(("xml_", "tu_"))]


def rule_params(name: str) -> list[dict]:
    """Danh sách tham số kèm nhãn/kiểu để dựng form."""
    out = []
    for p in inspect.signature(RULES[name]).parameters.values():
        if p.name == "path":
            continue
        label, kind, *rest = PARAM_UI_RULE.get((name, p.name)) or PARAM_UI.get(p.name, (p.name, "auto"))
        required = p.default is inspect.Parameter.empty
        from .i18n import tr
        out.append({"name": p.name, "label": tr(label), "kind": kind, "choices": rest[0] if rest else None,
                    "required": required, "default": None if required else p.default})
    return out


def parse_value(kind: str, text: str):
    """Chuyển chữ người dùng gõ trong form thành giá trị cho luật. Chuỗi rỗng → None."""
    text = (text or "").strip()
    if not text:
        return None
    if kind == "list":
        return [x.strip() for x in text.split(";") if x.strip()]
    if kind == "bool":
        if text not in BOOL_TEXT:
            raise ValueError("chọn Có hoặc Không")
        return BOOL_TEXT[text]
    if kind == "auto" and text[:1] in "{[":
        try:
            return json.loads(text)
        except ValueError:
            pass
    if kind in ("number", "auto"):
        for conv in (int, float):
            try:
                return conv(text.replace(",", ".") if conv is float else text)
            except ValueError:
                pass
        if kind == "number":
            raise ValueError(f"'{text}' không phải là số")
    return text


def format_value(kind: str, value) -> str:
    if value is None:
        return ""
    if kind == "list":
        return "; ".join(str(v) for v in _as_list(value))
    if kind == "bool":
        return "Có" if value else "Không"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


# ====================================================================== kiểm tra khai báo


def describe() -> str:
    lines = []
    for name, fn in RULES.items():
        params = [p for p in inspect.signature(fn).parameters.values() if p.name != "path"]
        sig = ", ".join(p.name if p.default is inspect.Parameter.empty else f"{p.name}={p.default!r}"
                        for p in params)
        lines.append(f"{name}({sig})\n    {(fn.__doc__ or '').strip()}")
    return "\n".join(lines)


def validate(spec: dict) -> str | None:
    """Trả về thông báo lỗi nếu khai báo luật sai tên / thiếu tham số."""
    name = spec.get("luat")
    if name not in RULES:
        return f"luật '{name}' không tồn tại"
    params = {k: v for k, v in spec.items() if k != "luat"}
    try:
        inspect.signature(RULES[name]).bind(Path("x"), **params)
    except TypeError as exc:
        return f"luật '{name}': {exc}"
    return None


def make_check(specs):
    """Tạo hàm chấm từ một luật hoặc danh sách luật (phải đúng tất cả)."""
    specs = _as_list(specs)

    def check(path: Path) -> bool | None:
        results = [RULES[s["luat"]](path, **{k: v for k, v in s.items() if k != "luat"}) for s in specs]
        if any(r is False or (r is not None and not r) for r in results):
            return False
        return None if any(r is None for r in results) else True

    return check
