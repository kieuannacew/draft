"""Nhập danh sách học viên từ file CSV hoặc Excel (.xlsx).

File cần một dòng tiêu đề. Tên cột nhận linh hoạt (có dấu / không dấu, hoa / thường):
    Họ tên (bắt buộc) · Tên đăng nhập · Mật khẩu · Lớp · Hạn dùng
Thiếu "Tên đăng nhập" → tự tạo từ họ tên; thiếu "Mật khẩu" → tạo ngẫu nhiên.
"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from .accounts import STUDENT, USERNAME_RE, AccountError, AccountStore, random_password, slugify

COLUMNS = {
    "username": ("tendangnhap", "taikhoan", "username", "user", "login", "tendn", "matk"),
    "ho_ten": ("hoten", "hovaten", "hotenhocvien", "hocvien", "fullname", "name", "ten"),
    "password": ("matkhau", "password", "pass", "mk"),
    "lop": ("lop", "class", "tenlop"),
    "han_dung": ("handung", "hethan", "ngayhethan", "expiry", "hansudung"),
}
TEMPLATE_HEADER = ["Họ tên", "Tên đăng nhập", "Mật khẩu", "Lớp", "Hạn dùng"]
TEMPLATE_ROWS = [["Nguyễn Văn An", "", "", "10A1", ""], ["Trần Thị Bình", "binhtt", "", "10A1", "2026-12-31"]]


class ImportFileError(ValueError):
    pass


@dataclass
class Entry:
    row: int                    # số dòng trong file (tính cả dòng tiêu đề)
    username: str
    ho_ten: str
    password: str
    lop: str = ""
    han_dung: str | None = None
    error: str = ""


def _key(header) -> str:
    return slugify(str(header or ""))


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read_table(path: Path) -> list[list[str]]:
    """Đọc toàn bộ ô của file thành danh sách dòng (chuỗi)."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        rows = [[_cell(v) for v in row] for row in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
        return rows
    if suffix in (".csv", ".txt"):
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1258", "cp1252"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [[_cell(v) for v in row] for row in csv.reader(text.splitlines(), dialect)]
    raise ImportFileError("Chỉ nhận file .csv hoặc .xlsx (Excel).")


def _find_header(rows: list[list[str]]) -> tuple[int, dict[str, int]]:
    for i, row in enumerate(rows[:10]):
        mapping = {}
        for col, header in enumerate(row):
            k = _key(header)
            for field, aliases in COLUMNS.items():
                if field not in mapping and k in aliases:
                    mapping[field] = col
        if "ho_ten" in mapping or "username" in mapping:
            return i, mapping
    raise ImportFileError("Không tìm thấy dòng tiêu đề. File cần có cột “Họ tên” (và tuỳ chọn: Tên đăng nhập, "
                       "Mật khẩu, Lớp, Hạn dùng). Bấm “Tải file mẫu” để xem ví dụ.")


def _date(text: str) -> str | None:
    if not text:
        return None
    m = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})", text)     # 31/12/2026
    if m:
        d, mth, y = map(int, m.groups())
        return date(y, mth, d).isoformat()
    return date.fromisoformat(text[:10]).isoformat()


def plan(path: Path, store: AccountStore) -> list[Entry]:
    """Đọc file và chuẩn bị danh sách tài khoản sẽ tạo (chưa ghi gì)."""
    rows = read_table(path)
    head, cols = _find_header(rows)
    get = lambda row, f: row[cols[f]].strip() if f in cols and cols[f] < len(row) else ""  # noqa: E731
    entries, taken = [], set()
    for i, row in enumerate(rows[head + 1:], start=head + 2):
        if not any(c.strip() for c in row):
            continue
        ho_ten, username = get(row, "ho_ten"), get(row, "username").lower()
        e = Entry(i, username, ho_ten, get(row, "password") or random_password(), get(row, "lop"))
        try:
            e.han_dung = _date(get(row, "han_dung"))
        except ValueError:
            e.error = "Hạn dùng không đúng dạng (YYYY-MM-DD hoặc DD/MM/YYYY)"
        if not ho_ten and not username:
            e.error = "Thiếu họ tên"
        elif not username:
            e.username = store.unique_username(ho_ten, taken)
        elif not USERNAME_RE.fullmatch(username):
            e.error = "Tên đăng nhập chỉ gồm chữ thường không dấu, số, _ hoặc . (3–32 ký tự)"
        elif username in store.accounts:
            e.error = "Tên đăng nhập đã tồn tại"
        elif username in taken:
            e.error = "Trùng tên đăng nhập với dòng khác trong file"
        if not e.ho_ten:
            e.ho_ten = e.username
        if len(e.password) < 4 and not e.error:
            e.error = "Mật khẩu phải có ít nhất 4 ký tự"
        taken.add(e.username)
        entries.append(e)
    if not entries:
        raise ImportFileError("File không có dòng học viên nào dưới dòng tiêu đề.")
    return entries


def apply(entries: list[Entry], store: AccountStore, giao_vien: str | None) -> tuple[list[Entry], list[str]]:
    """Tạo tài khoản học viên cho các dòng hợp lệ. Trả về (đã tạo, lỗi)."""
    created, errors = [], []
    for e in entries:
        if e.error:
            errors.append(f"Dòng {e.row}: {e.error}")
            continue
        try:
            store.create(e.username, e.ho_ten, e.password, STUDENT, han_dung=e.han_dung, lop=e.lop,
                         giao_vien=giao_vien)
            created.append(e)
        except AccountError as exc:
            errors.append(f"Dòng {e.row}: {exc}")
    return created, errors


def write_template(path: Path) -> None:
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(TEMPLATE_HEADER)
            w.writerows(TEMPLATE_ROWS)
        return
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "HocVien"
    ws.append(TEMPLATE_HEADER)
    for r in TEMPLATE_ROWS:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, width in zip("ABCDE", (24, 18, 14, 10, 14)):
        ws.column_dimensions[col].width = width
    ws["G1"] = "Chỉ cột Họ tên là bắt buộc. Để trống Tên đăng nhập / Mật khẩu để app tự tạo."
    wb.save(path)
