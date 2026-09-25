"""Đề luyện Word (kiểu MO-100)."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

from ..core import Exam, Project, Task
from ..ooxml import any_part_contains, norm, part

# ---------------------------------------------------------------- tiện ích


def _doc(path: Path):
    return Document(path)


def _find(doc, text: str):
    """Tìm đoạn văn có nội dung bắt đầu bằng `text` (bỏ qua hoa/thường)."""
    target = norm(text)
    for p in doc.paragraphs:
        if norm(p.text).startswith(target):
            return p
    return None


def _full_text(doc) -> str:
    chunks = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            chunks.extend(c.text for c in row.cells)
    return "\n".join(chunks)


def _plain(xml: str) -> str:
    return norm(re.sub(r"<[^>]+>", "", xml))


# ================================================================ Dự án 1

HEADINGS = ["Giới thiệu", "Kết quả kinh doanh", "Kế hoạch năm tới"]


def build_report(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Báo cáo thường niên")
    doc.add_paragraph(HEADINGS[0])
    doc.add_paragraph(
        "Công ty ABC được thành lập năm 2010, hoạt động trong lĩnh vực phân phối "
        "thiết bị văn phòng. Sau hơn 15 năm, công ty ABC đã có mặt tại 20 tỉnh thành.")
    doc.add_paragraph(HEADINGS[1])
    doc.add_paragraph(
        "Năm qua, doanh thu tăng 18% so với cùng kỳ. Lợi nhuận sau thuế đạt 12 tỷ đồng, "
        "vượt 5% kế hoạch đề ra.")
    doc.add_paragraph(HEADINGS[2])
    doc.add_paragraph(
        "Công ty ABC sẽ mở rộng sang thị trường thương mại điện tử và đầu tư hệ thống "
        "kho vận mới tại miền Trung.")
    doc.save(path)


def chk_title_style(path):
    p = _find(_doc(path), "Báo cáo thường niên")
    return p is not None and p.style.name == "Title"


def chk_headings(path):
    doc = _doc(path)
    for h in HEADINGS:
        p = _find(doc, h)
        if p is None or p.style.name != "Heading 1":
            return False
    return True


def chk_replace(path):
    text = norm(_full_text(_doc(path)))
    return "abc" not in text and text.count("công ty xyz") == 3


def chk_toc(path):
    return any_part_contains(path, r"word/document\.xml", r"(instrText[^>]*>\s*TOC\b|w:instr=\"\s*TOC\b)")


def chk_page_number(path):
    return any_part_contains(path, r"word/footer\d*\.xml",
                             r"(instrText[^>]*>\s*PAGE\b|w:instr=\"\s*PAGE\b)")


def chk_watermark(path):
    return any_part_contains(path, r"word/header\d*\.xml", r"BẢO MẬT", re.IGNORECASE)


def chk_landscape(path):
    s = _doc(path).sections[0]
    return s.orientation == WD_ORIENT.LANDSCAPE and s.page_width > s.page_height


# ================================================================ Dự án 2

BULLETS = ["Khai mạc", "Báo cáo chuyên đề", "Thảo luận và bế mạc"]
TAB_LINES = [
    "Thời gian\tNội dung\tNgười trình bày",
    "8:00\tKhai mạc\tBan tổ chức",
    "9:00\tChuyên đề\tTS. Nguyễn Văn A",
    "10:30\tThảo luận\tKhách mời",
]


def build_invite(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Thư mời hội thảo")
    doc.add_paragraph("Kính gửi: Quý đối tác")
    doc.add_paragraph("Chúng tôi trân trọng kính mời Quý vị tham dự hội thảo "
                      "“Chuyển đổi số trong doanh nghiệp vừa và nhỏ”.")
    doc.add_paragraph("Nội dung chương trình:")
    for b in BULLETS:
        doc.add_paragraph(b)
    doc.add_paragraph("Lịch trình chi tiết:")
    for line in TAB_LINES:
        doc.add_paragraph(line)
    doc.add_paragraph("Trân trọng kính mời.")
    doc.save(path)


def chk_center_title(path):
    p = _find(_doc(path), "Thư mời hội thảo")
    if p is None:
        return False
    align = p.alignment if p.alignment is not None else p.style.paragraph_format.alignment
    return align == WD_ALIGN_PARAGRAPH.CENTER


def chk_bullets(path):
    doc = _doc(path)
    for b in BULLETS:
        p = _find(doc, b)
        if p is None:
            return False
        has_num = p._p.pPr is not None and p._p.pPr.numPr is not None
        if not (has_num or p.style.name.startswith("List")):
            return False
    return True


def chk_table(path):
    for t in _doc(path).tables:
        if len(t.columns) == 3 and len(t.rows) >= 4 and norm(t.cell(0, 0).text) == norm("Thời gian"):
            return True
    return False


def chk_doc_title(path):
    return norm(_doc(path).core_properties.title or "") == norm("Thư mời hội thảo")


def chk_track_changes(path):
    settings = part(path, "word/settings.xml")
    m = re.search(r"<w:trackRevisions(?:\s+w:val=\"(\w+)\")?\s*/>", settings)
    return bool(m) and (m.group(1) or "true") not in ("0", "false", "off")


def chk_footnote(path):
    return "danh sách đối tác đính kèm" in _plain(part(path, "word/footnotes.xml"))


# ================================================================ Đề thi

EXAM = Exam(
    code="WORD",
    name="Microsoft Word (MO-100)",
    projects=[
        Project(
            name="Dự án 1 – Báo cáo thường niên",
            name_en='Project 1 – Annual report',
            intro_en="You need to finish the company's annual report before sending it to shareholders.",
            filename="BaoCao.docx",
            intro="Bạn cần hoàn thiện báo cáo thường niên của công ty trước khi gửi cổ đông.",
            build=build_report,
            tasks=[
                Task("Áp dụng kiểu (style) Title cho dòng đầu tiên “Báo cáo thường niên”.",
                     "Đặt con trỏ vào dòng đầu > Home > nhóm Styles > Title.",
                     chk_title_style,
                     'Apply the Title style to the first line “Báo cáo thường niên”.',
                     'Put the cursor in the first line > Home > Styles group > Title.'),
                Task("Áp dụng kiểu Heading 1 cho các dòng “Giới thiệu”, “Kết quả kinh doanh”, “Kế hoạch năm tới”.",
                     "Đặt con trỏ vào từng dòng > Home > Styles > Heading 1.",
                     chk_headings,
                     'Apply the Heading 1 style to the lines “Giới thiệu”, “Kết quả kinh doanh”, “Kế hoạch năm tới”.',
                     'Put the cursor in each line > Home > Styles > Heading 1.'),
                Task("Thay thế tất cả cụm từ “công ty ABC” thành “công ty XYZ”.",
                     "Home > Replace (Ctrl+H) > Find: ABC, Replace: XYZ > Replace All.",
                     chk_replace,
                     'Replace all instances of “công ty ABC” with “công ty XYZ”.',
                     'Home > Replace (Ctrl+H) > Find: ABC, Replace: XYZ > Replace All.'),
                Task("Chèn mục lục tự động (Table of Contents) vào đầu tài liệu.",
                     "Đặt con trỏ ở đầu tài liệu > References > Table of Contents > Automatic Table 1.",
                     chk_toc,
                     'Insert an automatic Table of Contents at the beginning of the document.',
                     'Put the cursor at the start of the document > References > Table of Contents > Automatic Table 1.'),
                Task("Chèn số trang vào chân trang (footer).",
                     "Insert > Page Number > Bottom of Page > Plain Number 2.",
                     chk_page_number,
                     'Insert page numbers in the footer.',
                     'Insert > Page Number > Bottom of Page > Plain Number 2.'),
                Task("Thêm watermark văn bản “BẢO MẬT”.",
                     "Design > Watermark > Custom Watermark > Text watermark > Text: BẢO MẬT > OK.",
                     chk_watermark,
                     'Add the text watermark “BẢO MẬT”.',
                     'Design > Watermark > Custom Watermark > Text watermark > Text: BẢO MẬT > OK.'),
                Task("Đổi hướng trang của tài liệu thành ngang (Landscape).",
                     "Layout > Orientation > Landscape.",
                     chk_landscape,
                     'Change the orientation of the document to Landscape.',
                     'Layout > Orientation > Landscape.'),
            ],
        ),
        Project(
            name="Dự án 2 – Thư mời hội thảo",
            name_en='Project 2 – Seminar invitation',
            intro_en='You are writing an invitation for guests to the company seminar.',
            filename="ThuMoi.docx",
            intro="Bạn soạn thư mời khách tham dự hội thảo của công ty.",
            build=build_invite,
            tasks=[
                Task("Căn giữa dòng tiêu đề “Thư mời hội thảo”.",
                     "Đặt con trỏ vào dòng tiêu đề > Home > Center (Ctrl+E).",
                     chk_center_title,
                     'Center the title line “Thư mời hội thảo”.',
                     'Put the cursor in the title line > Home > Center (Ctrl+E).'),
                Task("Định dạng 3 dòng “Khai mạc”, “Báo cáo chuyên đề”, “Thảo luận và bế mạc” thành danh sách dấu đầu dòng.",
                     "Chọn 3 dòng > Home > Bullets.",
                     chk_bullets,
                     'Format the 3 lines “Khai mạc”, “Báo cáo chuyên đề”, “Thảo luận và bế mạc” as a bulleted list.',
                     'Select the 3 lines > Home > Bullets.'),
                Task("Chuyển 4 dòng lịch trình (phân cách bằng Tab) thành bảng 3 cột.",
                     "Chọn 4 dòng từ “Thời gian…” đến “10:30…” > Insert > Table > Convert Text to Table > "
                     "Separate text at: Tabs > OK.",
                     chk_table,
                     'Convert the 4 schedule lines (separated by tabs) into a 3-column table.',
                     'Select the 4 lines from “Thời gian…” to “10:30…” > Insert > Table > Convert Text to Table > Separate text at: Tabs > OK.'),
                Task("Chèn chú thích cuối trang (footnote) sau “Quý đối tác” với nội dung “Danh sách đối tác đính kèm”.",
                     "Đặt con trỏ sau chữ “Quý đối tác” > References > Insert Footnote (Alt+Ctrl+F) > gõ nội dung.",
                     chk_footnote,
                     'Insert a footnote after “Quý đối tác” with the text “Danh sách đối tác đính kèm”.',
                     'Put the cursor after “Quý đối tác” > References > Insert Footnote (Alt+Ctrl+F) > type the text.'),
                Task("Đặt thuộc tính Title của tài liệu là “Thư mời hội thảo”.",
                     "File > Info > Properties (bên phải) > Title > gõ nội dung.",
                     chk_doc_title,
                     "Set the document's Title property to “Thư mời hội thảo”.",
                     'File > Info > Properties (on the right) > Title > type the text.'),
                Task("Bật chế độ theo dõi thay đổi (Track Changes).",
                     "Review > Track Changes (Ctrl+Shift+E).",
                     chk_track_changes,
                     'Turn on Track Changes.',
                     'Review > Track Changes (Ctrl+Shift+E).'),
            ],
        ),
    ],
)
