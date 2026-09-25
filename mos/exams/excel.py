"""Đề luyện Excel (kiểu MO-200)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.cell import range_boundaries

from ..core import Exam, Project, Task
from ..ooxml import any_part_contains, norm

# ---------------------------------------------------------------- tiện ích


def _wb(path: Path):
    return load_workbook(path)  # giữ công thức (data_only=False)


def _sheet(wb, preferred: str | None = None):
    if preferred and preferred in wb.sheetnames:
        return wb[preferred]
    return wb.worksheets[0]


def _formula(cell) -> str:
    v = cell.value
    text = getattr(v, "text", v)  # ArrayFormula có thuộc tính .text
    return str(text).upper().replace(" ", "") if isinstance(text, str) and text.startswith("=") else ""


def _defined_name_value(wb, name: str) -> str:
    """Hỗ trợ cả tên phạm vi Workbook lẫn phạm vi Sheet."""
    target = name.casefold()
    for key, dn in wb.defined_names.items():
        if key.casefold() == target:
            return dn.attr_text
    for ws in wb.worksheets:
        for key, dn in ws.defined_names.items():
            if key.casefold() == target:
                return dn.attr_text
    return ""


# ================================================================ Dự án 1

SALES = [
    ("Tháng 1", "Bút bi", 120, 5000),
    ("Tháng 1", "Vở", 80, 12000),
    ("Tháng 2", "Bút bi", 150, 5000),
    ("Tháng 2", "Thước kẻ", 60, 7000),
    ("Tháng 3", "Vở", 95, 12000),
    ("Tháng 3", "Bút chì", 200, 3000),
    ("Tháng 4", "Bút bi", 170, 5000),
    ("Tháng 4", "Tẩy", 90, 4000),
    ("Tháng 5", "Vở", 110, 12000),
    ("Tháng 5", "Thước kẻ", 75, 7000),
]


def build_sales(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Tháng", "Sản phẩm", "Số lượng", "Đơn giá", "Thành tiền"])
    for row in SALES:
        ws.append(list(row))
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, width in zip("ABCDE", (12, 14, 11, 11, 14)):
        ws.column_dimensions[col].width = width
    wb.save(path)


def chk_sheet_renamed(path):
    return "DoanhSo" in _wb(path).sheetnames and "Sheet1" not in _wb(path).sheetnames


def chk_amount_formula(path):
    ws = _sheet(_wb(path), "DoanhSo")
    for r in range(2, 12):
        f = _formula(ws[f"E{r}"])
        by_ref = f"C{r}" in f and f"D{r}" in f
        by_table = "[@[SỐLƯỢNG]]" in f or "[@SỐLƯỢNG]" in f
        if not (f and "*" in f and (by_ref or by_table)):
            return False
    return True


def chk_table(path):
    ws = _sheet(_wb(path), "DoanhSo")
    for table in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        if (min_col, min_row, max_col) == (1, 1, 5) and max_row >= 11:
            return True
    return False


def chk_freeze(path):
    return _sheet(_wb(path), "DoanhSo").freeze_panes == "A2"


def chk_number_format(path):
    ws = _sheet(_wb(path), "DoanhSo")
    return all(ws[f"{c}{r}"].number_format not in ("General", "@")
               for c in "DE" for r in range(2, 12))


def chk_total(path):
    ws = _sheet(_wb(path), "DoanhSo")
    f = _formula(ws["G2"])
    label_ok = norm(str(ws["G1"].value or "")) == norm("Tổng doanh thu")
    return label_ok and "SUM(" in f and ("E2:E11" in f or "[THÀNHTIỀN]" in f)


# ================================================================ Dự án 2

STAFF = [
    ("Nguyễn Văn An", "Kế toán", 9500000),
    ("Trần Thị Bình", "Kinh doanh", 12000000),
    ("Lê Văn Cường", "Kỹ thuật", 15000000),
    ("Phạm Thị Dung", "Kế toán", 11000000),
    ("Hoàng Văn Em", "Kinh doanh", 8000000),
    ("Vũ Thị Giang", "Nhân sự", 10000000),
    ("Đặng Văn Hùng", "Kỹ thuật", 13500000),
    ("Bùi Thị Lan", "Kế toán", 8500000),
    ("Đỗ Văn Minh", "Kinh doanh", 14000000),
    ("Ngô Thị Nga", "Nhân sự", 9000000),
]


def build_staff(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "NhanVien"
    ws.append(["Họ tên", "Phòng", "Lương", "Xếp loại"])
    for row in STAFF:
        ws.append(list(row))
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in range(2, 12):
        ws[f"C{r}"].number_format = "#,##0"
    for col, width in zip("ABCD", (18, 13, 13, 10)):
        ws.column_dimensions[col].width = width
    wb.save(path)


def chk_if(path):
    ws = _sheet(_wb(path), "NhanVien")
    for r in range(2, 12):
        f = _formula(ws[f"D{r}"])
        if not (f.startswith("=IF(") and f"C{r}" in f and '"CAO"' in f and '"THẤP"' in f):
            return False
    return True


def chk_sorted(path):
    ws = _sheet(_wb(path), "NhanVien")
    rows = [(ws[f"A{r}"].value, ws[f"C{r}"].value) for r in range(2, 12)]
    salaries = [s for _, s in rows]
    same_data = sorted(rows, key=str) == sorted(((n, s) for n, _, s in STAFF), key=str)
    return same_data and salaries == sorted(salaries, reverse=True)


def chk_conditional(path):
    ws = _sheet(_wb(path), "NhanVien")
    for cf in ws.conditional_formatting:
        if all(cell in cf.sqref for cell in ("C2", "C6", "C11")) and cf.rules:
            return True
    return False


def chk_named_range(path):
    value = _defined_name_value(_wb(path), "LuongNV").replace("'", "")
    return value.endswith("$C$2:$C$11")


def chk_countif(path):
    ws = _sheet(_wb(path), "NhanVien")
    f = _formula(ws["F2"])
    return "COUNTIF(" in f and "B2:B11" in f


def chk_chart(path):
    return any_part_contains(path, r"xl/charts/chart\d+\.xml", r'<(\w+:)?barDir val="col"')


def chk_landscape(path):
    return _sheet(_wb(path), "NhanVien").page_setup.orientation == "landscape"


# ================================================================ Đề thi

EXAM = Exam(
    code="EXCEL",
    name="Microsoft Excel (MO-200)",
    projects=[
        Project(
            name="Dự án 1 – Doanh số",
            name_en='Project 1 – Sales',
            intro_en='You manage the stationery sales sheet for the first 5 months of the year.',
            filename="DoanhSo.xlsx",
            intro="Bạn quản lý bảng doanh số văn phòng phẩm 5 tháng đầu năm.",
            build=build_sales,
            tasks=[
                Task("Đổi tên trang tính “Sheet1” thành “DoanhSo”.",
                     "Nhấp đúp vào tab Sheet1 ở dưới cùng (hoặc chuột phải > Rename), gõ DoanhSo, Enter.",
                     chk_sheet_renamed,
                     'Rename the worksheet “Sheet1” to “DoanhSo”.',
                     'Double-click the Sheet1 tab at the bottom (or right-click > Rename), type DoanhSo, press Enter.'),
                Task("Trong ô E2:E11, nhập công thức tính Thành tiền = Số lượng × Đơn giá.",
                     "Chọn E2, gõ =C2*D2, Enter. Kéo nút fill (góc phải dưới ô) xuống E11.",
                     chk_amount_formula,
                     'In cells E2:E11, enter a formula for Amount = Quantity × Unit price.',
                     'Select E2, type =C2*D2, press Enter. Drag the fill handle (bottom-right corner) down to E11.'),
                Task("Định dạng vùng A1:E11 thành bảng (Table) với kiểu bất kỳ, có hàng tiêu đề.",
                     "Chọn A1:E11 > Home > Format as Table > chọn kiểu > tick ‘My table has headers’ > OK (hoặc Ctrl+T).",
                     chk_table,
                     'Format the range A1:E11 as a table (any style) with a header row.',
                     'Select A1:E11 > Home > Format as Table > choose a style > check ‘My table has headers’ > OK (or Ctrl+T).'),
                Task("Cố định (freeze) hàng tiêu đề để luôn hiển thị khi cuộn.",
                     "View > Freeze Panes > Freeze Top Row.",
                     chk_freeze,
                     'Freeze the header row so it stays visible while scrolling.',
                     'View > Freeze Panes > Freeze Top Row.'),
                Task("Áp dụng định dạng số có dấu phân cách hàng nghìn hoặc tiền tệ cho D2:E11.",
                     "Chọn D2:E11 > Home > nhóm Number > Comma Style (,) hoặc Accounting/Currency.",
                     chk_number_format,
                     'Apply a number format with a thousands separator or currency to D2:E11.',
                     'Select D2:E11 > Home > Number group > Comma Style (,) or Accounting/Currency.'),
                Task("Ô G1 nhập “Tổng doanh thu”; ô G2 dùng hàm SUM tính tổng cột Thành tiền (E2:E11).",
                     "Gõ Tổng doanh thu vào G1. Chọn G2, gõ =SUM(E2:E11), Enter.",
                     chk_total,
                     'In cell G1 enter “Tổng doanh thu”; in G2 use SUM to total the Amount column (E2:E11).',
                     'Type Tổng doanh thu in G1. Select G2, type =SUM(E2:E11), press Enter.'),
            ],
        ),
        Project(
            name="Dự án 2 – Nhân viên",
            name_en='Project 2 – Staff',
            intro_en='You work in HR and need to complete the payroll sheet.',
            filename="NhanVien.xlsx",
            intro="Bạn là nhân viên phòng nhân sự, cần hoàn thiện bảng lương.",
            build=build_staff,
            tasks=[
                Task("Ô D2:D11 dùng hàm IF: nếu Lương ≥ 10.000.000 thì “Cao”, ngược lại “Thấp”.",
                     'Chọn D2, gõ =IF(C2>=10000000,"Cao","Thấp"), Enter, kéo xuống D11.',
                     chk_if,
                     'In D2:D11 use IF: if Salary ≥ 10,000,000 then “Cao”, otherwise “Thấp”.',
                     'Select D2, type =IF(C2>=10000000,"Cao","Thấp"), press Enter, drag down to D11.'),
                Task("Sắp xếp dữ liệu (A1:D11) theo cột Lương giảm dần.",
                     "Chọn một ô trong bảng > Data > Sort > Sort by: Lương, Order: Largest to Smallest > OK.",
                     chk_sorted,
                     'Sort the data (A1:D11) by the Salary column, largest to smallest.',
                     'Select a cell in the table > Data > Sort > Sort by: Lương, Order: Largest to Smallest > OK.'),
                Task("Áp dụng định dạng có điều kiện (quy tắc bất kỳ) cho vùng C2:C11.",
                     "Chọn C2:C11 > Home > Conditional Formatting > ví dụ Data Bars hoặc Highlight Cells Rules.",
                     chk_conditional,
                     'Apply conditional formatting (any rule) to C2:C11.',
                     'Select C2:C11 > Home > Conditional Formatting > e.g. Data Bars or Highlight Cells Rules.'),
                Task("Đặt tên vùng C2:C11 là “LuongNV”.",
                     "Chọn C2:C11, gõ LuongNV vào Name Box (bên trái thanh công thức), Enter. "
                     "Hoặc Formulas > Define Name.",
                     chk_named_range,
                     'Name the range C2:C11 “LuongNV”.',
                     'Select C2:C11, type LuongNV in the Name Box (left of the formula bar), press Enter. Or Formulas > Define Name.'),
                Task("Ô F2 dùng hàm COUNTIF đếm số nhân viên phòng “Kế toán” (vùng B2:B11).",
                     'Chọn F2, gõ =COUNTIF(B2:B11,"Kế toán"), Enter.',
                     chk_countif,
                     'In F2 use COUNTIF to count staff in the “Kế toán” department (range B2:B11).',
                     'Select F2, type =COUNTIF(B2:B11,"Kế toán"), press Enter.'),
                Task("Chèn biểu đồ cột (Clustered Column) thể hiện Lương theo Họ tên.",
                     "Chọn A1:A11, giữ Ctrl chọn C1:C11 > Insert > Column Chart > Clustered Column.",
                     chk_chart,
                     'Insert a Clustered Column chart showing Salary by Name.',
                     'Select A1:A11, hold Ctrl and select C1:C11 > Insert > Column Chart > Clustered Column.'),
                Task("Đặt hướng trang in của trang tính là ngang (Landscape).",
                     "Page Layout > Orientation > Landscape.",
                     chk_landscape,
                     "Set the worksheet's print orientation to Landscape.",
                     'Page Layout > Orientation > Landscape.'),
            ],
        ),
    ],
)
