"""Đề luyện PowerPoint (kiểu MO-300)."""
from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from ..core import Exam, Project, Task
from ..ooxml import norm, part

# ---------------------------------------------------------------- tiện ích


def _prs(path: Path):
    return Presentation(path)


def _title(slide) -> str:
    t = slide.shapes.title
    return t.text_frame.text if t is not None and t.has_text_frame else ""


def _slide(prs, title: str):
    for s in prs.slides:
        if norm(_title(s)) == norm(title):
            return s
    return None


def _xml(slide) -> str:
    return slide._element.xml


def _new(slides_spec: list[tuple[int, str, str]], path: Path, widescreen=True) -> None:
    """slides_spec: (chỉ số layout, tiêu đề, nội dung)."""
    prs = Presentation()
    if widescreen:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for layout, title, body in slides_spec:
        slide = prs.slides.add_slide(prs.slide_layouts[layout])
        slide.shapes.title.text = title
        if body and len(slide.placeholders) > 1:
            slide.placeholders[1].text = body
    prs.save(path)


# ================================================================ Dự án 1


def build_product(path: Path) -> None:
    _new([
        (0, "Giới thiệu sản phẩm", "Năm 2026"),
        (1, "Tính năng nổi bật", "Nhỏ gọn\nPin 48 giờ\nChống nước"),
        (1, "Giá bán", "Bản tiêu chuẩn: 2.990.000đ\nBản cao cấp: 3.990.000đ"),
        (1, "Liên hệ", "Hotline: 1900 0000\nEmail: sales@example.com"),
    ], path)


def chk_notes(path):
    s = _prs(path).slides[0]
    return s.has_notes_slide and "chào mừng khán giả" in norm(s.notes_slide.notes_text_frame.text)


def chk_transitions(path):
    return all("<p:transition" in _xml(s) for s in _prs(path).slides)


def chk_hidden(path):
    s = _slide(_prs(path), "Giá bán")
    return s is not None and s._element.get("show") in ("0", "false")


def chk_animation(path):
    s = _slide(_prs(path), "Giới thiệu sản phẩm")
    return s is not None and "<p:timing" in _xml(s) and 'presetClass="entr"' in _xml(s)


def chk_size(path):
    prs = _prs(path)
    return abs(prs.slide_width / prs.slide_height - 4 / 3) < 0.01


def chk_new_slide(path):
    prs = _prs(path)
    last = prs.slides[len(prs.slides) - 1]
    return norm(_title(last)) == norm("Kế hoạch") and last.slide_layout.name == "Title and Content"


# ================================================================ Dự án 2


def build_quarter(path: Path) -> None:
    _new([
        (0, "Báo cáo quý 3", "Phòng Kinh doanh"),
        (5, "Doanh thu theo khu vực", ""),
        (5, "Quy trình bán hàng", ""),
        (5, "Xóa slide này", ""),
        (5, "Cảm ơn", ""),
    ], path)


def chk_table(path):
    s = _slide(_prs(path), "Doanh thu theo khu vực")
    if s is None:
        return False
    for sh in s.shapes:
        if sh.has_table and len(sh.table.columns) == 3 and len(sh.table.rows) == 4:
            return True
    return False


def chk_smartart(path):
    s = _slide(_prs(path), "Quy trình bán hàng")
    return s is not None and "drawingml/2006/diagram" in _xml(s)


def chk_deleted(path):
    return _slide(_prs(path), "Xóa slide này") is None and len(_prs(path).slides) == 4


def chk_section(path):
    xml = part(path, "ppt/presentation.xml")
    return re.search(r'<p14:section [^>]*name="Mở đầu"', xml) is not None


def chk_slide_numbers(path):
    slides = list(_prs(path).slides)[1:]  # cho phép ẩn trên slide tiêu đề
    return bool(slides) and all(re.search(r'<p:ph [^>]*type="sldNum"', _xml(s)) for s in slides)


# ================================================================ Đề thi

EXAM = Exam(
    code="POWERPOINT",
    name="Microsoft PowerPoint (MO-300)",
    projects=[
        Project(
            name="Dự án 1 – Giới thiệu sản phẩm",
            name_en='Project 1 – Product launch',
            intro_en='You are preparing a presentation introducing a new smartwatch.',
            filename="GioiThieu.pptx",
            intro="Bạn chuẩn bị bài thuyết trình giới thiệu đồng hồ thông minh mới.",
            build=build_product,
            tasks=[
                Task("Thêm ghi chú (Notes) “Chào mừng khán giả” cho slide 1.",
                     "Chọn slide 1 > View > Notes (hoặc bấm Notes dưới cùng) > gõ nội dung vào khung ghi chú.",
                     chk_notes,
                     'Add the speaker note “Chào mừng khán giả” to slide 1.',
                     'Select slide 1 > View > Notes (or click Notes at the bottom) > type in the notes pane.', chapter=2),
                Task("Áp dụng hiệu ứng chuyển trang (Transition) bất kỳ cho tất cả các slide.",
                     "Transitions > chọn hiệu ứng (vd Fade) > Apply To All.",
                     chk_transitions,
                     'Apply any transition to all slides.',
                     'Transitions > choose an effect (e.g. Fade) > Apply To All.', chapter=5),
                Task("Ẩn slide “Giá bán” khi trình chiếu.",
                     "Chuột phải vào slide “Giá bán” ở khung bên trái > Hide Slide.",
                     chk_hidden,
                     'Hide the slide “Giá bán” during the slide show.',
                     'Right-click the “Giá bán” slide in the left pane > Hide Slide.', chapter=2),
                Task("Thêm hiệu ứng xuất hiện (Entrance animation) bất kỳ cho một đối tượng trên slide 1.",
                     "Chọn tiêu đề slide 1 > Animations > chọn hiệu ứng nhóm Entrance (vd Fade, Fly In).",
                     chk_animation,
                     'Add any Entrance animation to an object on slide 1.',
                     'Select the slide 1 title > Animations > choose an Entrance effect (e.g. Fade, Fly In).', chapter=5),
                Task("Đổi kích thước slide sang tỉ lệ 4:3 (Standard).",
                     "Design > Slide Size > Standard (4:3) > Ensure Fit.",
                     chk_size,
                     'Change the slide size to 4:3 (Standard).',
                     'Design > Slide Size > Standard (4:3) > Ensure Fit.', chapter=1),
                Task("Thêm slide mới bố cục “Title and Content” ở cuối, tiêu đề “Kế hoạch”.",
                     "Chọn slide cuối > Home > New Slide > Title and Content > gõ tiêu đề Kế hoạch.",
                     chk_new_slide,
                     'Add a new “Title and Content” slide at the end with the title “Kế hoạch”.',
                     'Select the last slide > Home > New Slide > Title and Content > type the title Kế hoạch.', chapter=2),
            ],
        ),
        Project(
            name="Dự án 2 – Báo cáo quý",
            name_en='Project 2 – Quarterly report',
            intro_en='You are finishing the Q3 business results presentation.',
            filename="BaoCaoQuy.pptx",
            intro="Bạn hoàn thiện bài báo cáo kết quả kinh doanh quý 3.",
            build=build_quarter,
            tasks=[
                Task("Chèn bảng 3 cột, 4 hàng vào slide “Doanh thu theo khu vực”.",
                     "Chọn slide > Insert > Table > kéo chọn 3x4.",
                     chk_table,
                     'Insert a table with 3 columns and 4 rows on the slide “Doanh thu theo khu vực”.',
                     'Select the slide > Insert > Table > drag to select 3x4.', chapter=4),
                Task("Chèn một SmartArt bất kỳ vào slide “Quy trình bán hàng”.",
                     "Chọn slide > Insert > SmartArt > chọn kiểu (vd Process > Basic Process) > OK.",
                     chk_smartart,
                     'Insert any SmartArt graphic on the slide “Quy trình bán hàng”.',
                     'Select the slide > Insert > SmartArt > choose a layout (e.g. Process > Basic Process) > OK.', chapter=4),
                Task("Xóa slide “Xóa slide này”.",
                     "Chuột phải vào slide ở khung bên trái > Delete Slide.",
                     chk_deleted,
                     'Delete the slide “Xóa slide này”.',
                     'Right-click the slide in the left pane > Delete Slide.', chapter=2),
                Task("Tạo section tên “Mở đầu” bắt đầu từ slide 1.",
                     "Chuột phải vào slide 1 > Add Section > gõ Mở đầu > Rename.",
                     chk_section,
                     'Create a section named “Mở đầu” starting at slide 1.',
                     'Right-click slide 1 > Add Section > type Mở đầu > Rename.', chapter=2),
                Task("Hiển thị số trang (Slide number) trên tất cả các slide (có thể trừ slide tiêu đề).",
                     "Insert > Header & Footer > tick Slide number > Apply to All.",
                     chk_slide_numbers,
                     'Show slide numbers on all slides (the title slide may be excluded).',
                     'Insert > Header & Footer > check Slide number > Apply to All.', chapter=2),
            ],
        ),
    ],
)
