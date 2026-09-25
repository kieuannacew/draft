"""Lớp học trực tuyến (Supabase): đăng nhập / đăng ký, lớp và mã lớp, gửi kết quả + tiến độ học lên máy chủ.

Cấu hình máy chủ (một trong hai chỗ, chỗ trước được ưu tiên):
  * `cau_hinh.json` cạnh run.bat / .exe:   {"may_chu": {"url": "https://xxxx.supabase.co", "key": "<anon key>"}}
    → giáo viên gửi kèm file này trong bộ cài, học sinh mở app là tự kết nối.
  * `~/MOS_Practice/may_chu.json` (lưu từ nút "Máy chủ lớp học" ở màn đăng nhập).
Không có cấu hình → app chạy ngoại tuyến như cũ (tài khoản trên từng máy).

Tài khoản trên máy chủ: tên đăng nhập `abc` ↔ email nội bộ `abc@<EMAIL_DOMAIN>` (không gửi thư).
Kết quả / tiến độ luôn được ghi vào hàng đợi `DATA_DIR/cho_dong_bo.json` trước rồi gửi dần (mất mạng không mất dữ
liệu, có mạng lại tự gửi). Mọi hàm mạng chạy đồng bộ; giao diện gọi trong luồng phụ khi cần.
"""
from __future__ import annotations

import json
import os
import threading
import urllib.error
import urllib.parse
import urllib.request
import uuid
from datetime import datetime
from pathlib import Path

from .core import APP_DIR, DATA_DIR, app_dir

EMAIL_DOMAIN = "hocvien.mos"
MIN_PASSWORD = 6
PAGE = 1000                      # Supabase trả tối đa 1000 dòng mỗi lần
ROLES = ("quan_tri", "giao_vien", "hoc_vien")


class CloudError(Exception):
    """Lỗi từ máy chủ, thông báo đã dịch sang tiếng Việt (khóa tr)."""


class OfflineError(CloudError):
    """Không kết nối được máy chủ."""


# ====================================================================== cấu hình


def _local_config() -> Path:
    return APP_DIR / "may_chu.json"


def _read_server(data: dict) -> dict | None:
    s = data.get("may_chu") if isinstance(data.get("may_chu"), dict) else data
    url, key = str(s.get("url") or "").strip().rstrip("/"), str(s.get("key") or "").strip()
    return {"url": url, "key": key} if url and key else None


def load_server() -> dict | None:
    """{"url", "key"} của máy chủ lớp học, hoặc None (chạy ngoại tuyến)."""
    for path in (app_dir() / "cau_hinh.json", _local_config()):
        try:
            data = json.loads(path.read_text(encoding="utf-8-sig"))
        except (OSError, ValueError):
            continue
        if isinstance(data, dict):
            server = _read_server(data)
            if server:
                return server
    return None


def server_source() -> Path | None:
    """File đang chứa cấu hình máy chủ (để hiện cho quản trị biết)."""
    for path in (app_dir() / "cau_hinh.json", _local_config()):
        try:
            if _read_server(json.loads(path.read_text(encoding="utf-8-sig"))):
                return path
        except (OSError, ValueError, AttributeError):
            continue
    return None


def save_server(url: str, key: str) -> Path:
    """Lưu cấu hình vào cau_hinh.json cạnh app (để gửi kèm cho học sinh) và một bản trên máy
    (~/MOS_Practice/may_chu.json – tải bản app mới về vẫn nhớ máy chủ). Trả file cạnh app nếu ghi được."""
    url = url.strip().rstrip("/")
    if not url.startswith(("https://", "http://")):
        url = "https://" + url
    try:
        _local_config().parent.mkdir(parents=True, exist_ok=True)
        _local_config().write_text(json.dumps({"may_chu": {"url": url, "key": key.strip()}}, indent=2),
                                   encoding="utf-8")
    except OSError:
        pass
    target = app_dir() / "cau_hinh.json"
    try:
        data = json.loads(target.read_text(encoding="utf-8-sig"))
        if not isinstance(data, dict):
            data = {}
    except (OSError, ValueError):
        data = {}
    data["may_chu"] = {"url": url, "key": key.strip()}
    try:
        target.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        return target
    except OSError:
        return _local_config()


def clear_server() -> None:
    """Tắt máy chủ lớp học (xóa cấu hình ở cả hai chỗ, giữ các mục khác trong cau_hinh.json)."""
    target = app_dir() / "cau_hinh.json"
    try:
        data = json.loads(target.read_text(encoding="utf-8-sig"))
        if isinstance(data, dict) and "may_chu" in data:
            del data["may_chu"]
            target.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    except (OSError, ValueError):
        pass
    try:
        _local_config().unlink()
    except OSError:
        pass


# ====================================================================== lỗi máy chủ → tiếng Việt

_ERRORS = {
    "invalid_credentials": "Sai tên đăng nhập hoặc mật khẩu.",
    "invalid login credentials": "Sai tên đăng nhập hoặc mật khẩu.",
    "user_already_exists": "Tên đăng nhập này đã có người dùng.",
    "user already registered": "Tên đăng nhập này đã có người dùng.",
    "weak_password": "Mật khẩu phải có ít nhất 6 ký tự.",
    "email_not_confirmed": "Máy chủ đang bật xác nhận email. Quản trị hãy tắt “Confirm email” trong Supabase.",
    "signup_disabled": "Máy chủ đang tắt đăng ký tài khoản mới.",
    "over_request_rate_limit": "Thao tác quá nhanh, hãy thử lại sau ít phút.",
    "over_email_send_rate_limit": "Máy chủ đang bật xác nhận email. Quản trị hãy tắt “Confirm email” trong Supabase.",
    "sai_ma_lop": "Mã lớp không đúng.",
    "khong_du_quyen": "Bạn không có quyền làm việc này.",
    "quan_tri_cuoi": "Phải còn ít nhất một tài khoản quản trị.",
    "khong_tu_khoa": "Không thể tự khóa tài khoản của mình.",
    "khong_tu_xoa": "Không thể tự xóa tài khoản của mình.",
    "mat_khau_ngan": "Mật khẩu phải có ít nhất 6 ký tự.",
    "chua_dang_nhap": "Tài khoản bị khóa hoặc chưa đăng nhập.",
    "refresh_token_not_found": "Phiên đăng nhập đã hết hạn, hãy đăng nhập lại.",
    "jwt expired": "Phiên đăng nhập đã hết hạn, hãy đăng nhập lại.",
    "no api key found": "Khóa máy chủ (anon key) không đúng.",
    "invalid api key": "Khóa máy chủ (anon key) không đúng.",
    "pgrst202": "Máy chủ chưa được cài đặt: hãy chạy file may_chu/supabase_lop_hoc.sql trong Supabase.",
    "pgrst205": "Máy chủ chưa được cài đặt: hãy chạy file may_chu/supabase_lop_hoc.sql trong Supabase.",
    "42p01": "Máy chủ chưa được cài đặt: hãy chạy file may_chu/supabase_lop_hoc.sql trong Supabase.",
}


def _friendly(status: int, data) -> CloudError:
    texts = []
    if isinstance(data, dict):
        for k in ("error_code", "code", "msg", "message", "error_description", "error"):
            if data.get(k):
                texts.append(str(data[k]))
    elif data:
        texts.append(str(data))
    blob = " ".join(texts).lower()
    for key, msg in _ERRORS.items():
        if key in blob:
            return CloudError(msg)
    detail = next((t for t in texts if not t.isdigit()), "") or f"HTTP {status}"
    return CloudError("Máy chủ báo lỗi: " + detail[:200])


# ====================================================================== client


def email_of(username: str) -> str:
    return f"{username.strip().lower()}@{EMAIL_DOMAIN}"


class Cloud:
    """Kết nối tới Supabase của lớp học (một phiên đăng nhập)."""

    def __init__(self, url: str, key: str, timeout: float = 15):
        self.url, self.key, self.timeout = url.rstrip("/"), key, timeout
        self.access_token = ""
        self.refresh_token = ""
        self.user_id = ""
        self.profile: dict = {}
        self._lock = threading.Lock()

    @classmethod
    def from_config(cls) -> "Cloud | None":
        server = load_server()
        return cls(server["url"], server["key"]) if server else None

    # ------------------------------------------------------------ HTTP
    def _http(self, method: str, path: str, body=None, params: dict | None = None,
              headers: dict | None = None, auth: bool = True, retry: bool = True):
        url = self.url + path
        if params:
            url += "?" + urllib.parse.urlencode(params, safe="(),.:*")
        req = urllib.request.Request(url, method=method,
                                     data=None if body is None else json.dumps(body).encode("utf-8"))
        req.add_header("apikey", self.key)
        if auth and self.access_token:
            req.add_header("Authorization", "Bearer " + self.access_token)
        elif self.key.count(".") == 2:          # khóa "anon" kiểu cũ là JWT; khóa "sb_publishable_…" thì không gửi
            req.add_header("Authorization", "Bearer " + self.key)
        req.add_header("Content-Type", "application/json")
        req.add_header("Accept", "application/json")
        for k, v in (headers or {}).items():
            req.add_header(k, v)
        try:
            with urllib.request.urlopen(req, timeout=self.timeout) as r:
                raw = r.read()
                return json.loads(raw) if raw.strip() else None
        except urllib.error.HTTPError as e:
            raw = e.read()
            try:
                data = json.loads(raw)
            except ValueError:
                data = raw.decode("utf-8", "replace")
            expired = e.code == 401 and "jwt" in json.dumps(data).lower()
            if expired and retry and auth and self.refresh_token:
                self.refresh()
                return self._http(method, path, body, params, headers, auth, retry=False)
            raise _friendly(e.code, data) from None
        except (urllib.error.URLError, OSError, TimeoutError) as e:
            raise OfflineError("Không kết nối được máy chủ lớp học. Kiểm tra mạng Internet.") from e

    def _session(self, data: dict) -> None:
        if not data or not data.get("access_token"):
            raise CloudError("Máy chủ đang bật xác nhận email. Quản trị hãy tắt “Confirm email” trong Supabase.")
        self.access_token = data["access_token"]
        self.refresh_token = data.get("refresh_token", "")
        self.user_id = (data.get("user") or {}).get("id", self.user_id)

    def check(self) -> dict:
        """Kiểm tra máy chủ: khóa đúng chưa, đã chạy file SQL chưa, có đang bắt xác nhận email không."""
        settings = self._http("GET", "/auth/v1/settings", auth=False) or {}
        try:
            self._http("POST", "/rest/v1/rpc/lop_cua_toi", {}, auth=False)
            schema = True
        except OfflineError:
            raise
        except CloudError as exc:
            schema = "supabase_lop_hoc.sql" not in str(exc)
        return {"schema": schema, "autoconfirm": settings.get("mailer_autoconfirm", True),
                "signup": not settings.get("disable_signup", False)}

    # ------------------------------------------------------------ tài khoản
    def signup(self, username: str, password: str, ho_ten: str) -> dict:
        username = username.strip().lower()
        data = self._http("POST", "/auth/v1/signup", {"email": email_of(username), "password": password,
                                                      "data": {"username": username, "ho_ten": ho_ten.strip()}},
                          auth=False)
        self._session(data)
        return self.load_profile()

    def login(self, username: str, password: str) -> dict:
        data = self._http("POST", "/auth/v1/token", {"email": email_of(username), "password": password},
                          params={"grant_type": "password"}, auth=False)
        self._session(data)
        return self.load_profile()

    def refresh(self) -> None:
        with self._lock:
            data = self._http("POST", "/auth/v1/token", {"refresh_token": self.refresh_token},
                              params={"grant_type": "refresh_token"}, auth=False, retry=False)
            self._session(data)

    def logout(self) -> None:
        self.access_token = self.refresh_token = self.user_id = ""
        self.profile = {}

    def load_profile(self) -> dict:
        rows = self._http("GET", "/rest/v1/ho_so", params={"id": f"eq.{self.user_id}", "select": "*"})
        if not rows:
            raise CloudError("Máy chủ chưa được cài đặt: hãy chạy file may_chu/supabase_lop_hoc.sql trong Supabase.")
        self.profile = rows[0]
        if self.profile.get("khoa"):
            self.logout()
            raise CloudError("Tài khoản đã bị khóa. Liên hệ giáo viên.")
        return self.profile

    @property
    def role(self) -> str:
        return self.profile.get("vai_tro", "hoc_vien")

    def change_password(self, password: str) -> None:
        self._http("PUT", "/auth/v1/user", {"password": password})

    def rename_me(self, ho_ten: str) -> None:
        self.rpc("doi_ho_ten", p_ho_ten=ho_ten)
        self.profile["ho_ten"] = ho_ten.strip()

    # ------------------------------------------------------------ REST
    def rpc(self, name: str, **args):
        return self._http("POST", f"/rest/v1/rpc/{name}", args)

    def select(self, table: str, **params) -> list[dict]:
        """Đọc hết các dòng (tự chia trang 1000 dòng)."""
        out, offset = [], 0
        while True:
            rows = self._http("GET", f"/rest/v1/{table}", params={**params, "limit": PAGE, "offset": offset}) or []
            out += rows
            if len(rows) < PAGE:
                return out
            offset += PAGE

    def _in(self, ids) -> str:
        return "in.(" + ",".join(ids) + ")"

    # ------------------------------------------------------------ lớp học
    def my_classes(self) -> list[dict]:
        return self.rpc("lop_cua_toi") or []

    def create_class(self, ten: str) -> dict:
        return self.rpc("tao_lop", ten_lop=ten)

    def join_class(self, ma: str) -> dict:
        return self.rpc("vao_lop", ma_lop=ma)

    def rename_class(self, lop_id: str, ten: str) -> None:
        self._http("PATCH", "/rest/v1/lop", {"ten": ten.strip()}, params={"id": f"eq.{lop_id}"})

    def delete_class(self, lop_id: str) -> None:
        self._http("DELETE", "/rest/v1/lop", params={"id": f"eq.{lop_id}"})

    def leave_class(self, lop_id: str, hoc_vien: str | None = None) -> None:
        self._http("DELETE", "/rest/v1/thanh_vien",
                   params={"lop": f"eq.{lop_id}", "hoc_vien": f"eq.{hoc_vien or self.user_id}"})

    def members(self, lop_id: str) -> list[dict]:
        """Học sinh trong lớp: [{id, username, ho_ten, khoa, vao_luc}]."""
        rows = self.select("thanh_vien", lop=f"eq.{lop_id}", select="vao_luc,ho_so(id,username,ho_ten,khoa)",
                           order="vao_luc")
        return [{**(r.get("ho_so") or {}), "vao_luc": r.get("vao_luc", "")} for r in rows if r.get("ho_so")]

    def results(self, student_ids: list[str] | None = None) -> list[dict]:
        """Kết quả (mới nhất trước) của các học sinh, hoặc của chính mình."""
        ids = student_ids if student_ids is not None else [self.user_id]
        if not ids:
            return []
        out = []
        for i in range(0, len(ids), 50):
            out += self.select("ket_qua", hoc_vien=self._in(ids[i:i + 50]), order="lam_luc.desc", select="*")
        return sorted(out, key=lambda r: r.get("lam_luc", ""), reverse=True)

    def progress(self, student_ids: list[str] | None = None) -> list[dict]:
        ids = student_ids if student_ids is not None else [self.user_id]
        out = []
        for i in range(0, len(ids), 50):
            out += self.select("tien_do", hoc_vien=self._in(ids[i:i + 50]), select="*")
        return out

    # ------------------------------------------------------------ quản trị
    def profiles(self) -> list[dict]:
        return self.select("ho_so", select="*", order="vai_tro,username")

    def set_role(self, user_id: str, vai_tro: str) -> None:
        self.rpc("dat_vai_tro", p_id=user_id, p_vai_tro=vai_tro)

    def set_locked(self, user_id: str, khoa: bool) -> None:
        self.rpc("dat_khoa", p_id=user_id, p_khoa=khoa)

    def reset_password(self, user_id: str, password: str) -> None:
        self.rpc("dat_lai_mat_khau", p_id=user_id, p_mat_khau=password)

    def delete_account(self, user_id: str) -> None:
        self.rpc("xoa_tai_khoan", p_id=user_id)

    # ------------------------------------------------------------ gửi dữ liệu
    def send_result(self, data: dict) -> None:
        self._http("POST", "/rest/v1/ket_qua", [data], params={"on_conflict": "hoc_vien,ma_client"},
                   headers={"Prefer": "resolution=ignore-duplicates,return=minimal"})

    def send_progress(self, data: dict) -> None:
        self.rpc("luu_tien_do", p_bai=data["bai"], p_ten_bai=data.get("ten_bai", ""), p_xem=int(data.get("xem", 0)),
                 p_tong=int(data.get("tong", 0)), p_xong=bool(data.get("xong")), p_diem=data.get("diem", ""))


# ====================================================================== hàng đợi gửi dữ liệu

_QUEUE_LOCK = threading.Lock()
_flushing: set[str] = set()


def _queue_file() -> Path:
    return DATA_DIR / "cho_dong_bo.json"


def _read_queue() -> list[dict]:
    try:
        data = json.loads(_queue_file().read_text(encoding="utf-8"))
        return data if isinstance(data, list) else []
    except (OSError, ValueError):
        return []


def _write_queue(items: list[dict]) -> None:
    path = _queue_file()
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(items, ensure_ascii=False, indent=1), encoding="utf-8")
    os.replace(tmp, path)


def _iso(time_text: str) -> str:
    try:
        return datetime.strptime(time_text, "%Y-%m-%d %H:%M").astimezone().isoformat()
    except (TypeError, ValueError):
        return datetime.now().astimezone().isoformat()


def result_payload(entry: dict) -> dict:
    """Mục lịch sử (history.json) → dòng bảng ket_qua."""
    return {"ma_client": entry.get("ma") or uuid.uuid4().hex, "exam": entry.get("exam", ""),
            "exam_en": entry.get("exam_en") or "", "code": entry.get("code") or "", "mode": entry.get("mode", ""),
            "score": int(entry.get("score", 0)), "correct": int(entry.get("correct", 0)),
            "total": int(entry.get("total", 0)), "passed": bool(entry.get("passed")),
            "duration": int(entry.get("duration") or 0), "wrong": list(entry.get("wrong") or [])[:60],
            "lam_luc": _iso(entry.get("time", ""))}


def queue_result(user: str, entry: dict) -> None:
    with _QUEUE_LOCK:
        items = _read_queue()
        items.append({"user": user, "loai": "ket_qua", "data": result_payload(entry)})
        _write_queue(items)


def queue_progress(user: str, bai: str, ten_bai: str, xem: int, tong: int, xong: bool, diem: str = "") -> None:
    """Tiến độ bài giảng; gộp với mục cũ cùng bài (giữ số lớn nhất)."""
    with _QUEUE_LOCK:
        items = _read_queue()
        for it in items:
            if it["user"] == user and it["loai"] == "tien_do" and it["data"]["bai"] == bai:
                d = it["data"]
                d.update(ten_bai=ten_bai, tong=tong, xem=max(d.get("xem", 0), xem), xong=d.get("xong") or xong,
                         diem=diem or d.get("diem", ""))
                break
        else:
            items.append({"user": user, "loai": "tien_do", "data": {"bai": bai, "ten_bai": ten_bai, "xem": xem,
                                                                    "tong": tong, "xong": xong, "diem": diem}})
        _write_queue(items)


def pending(user: str) -> int:
    return sum(it.get("user") == user for it in _read_queue())


def flush(cloud: Cloud, user: str) -> int:
    """Gửi các mục đang chờ của `user`. Trả số mục đã gửi; mất mạng thì dừng, giữ phần còn lại."""
    with _QUEUE_LOCK:
        mine = [it for it in _read_queue() if it.get("user") == user]
    sent = []
    for it in mine:
        try:
            if it["loai"] == "ket_qua":
                cloud.send_result(it["data"])
            else:
                cloud.send_progress(it["data"])
        except OfflineError:
            break
        except CloudError:
            continue                 # dữ liệu hỏng / không có quyền: giữ lại, lần sau thử tiếp
        sent.append(it)
    if sent:
        with _QUEUE_LOCK:
            rest = _read_queue()
            for it in sent:
                if it in rest:
                    rest.remove(it)
            _write_queue(rest)
    return len(sent)


def flush_in_background(cloud: Cloud | None, user: str, done=None) -> None:
    """Gửi hàng đợi trong luồng phụ (không làm đứng giao diện). `done(n)` được gọi trong luồng phụ."""
    if cloud is None or not cloud.access_token or user in _flushing:
        return
    _flushing.add(user)

    def run():
        try:
            n = flush(cloud, user)
        except Exception:        # không để luồng phụ làm hỏng app
            n = 0
        finally:
            _flushing.discard(user)
        if done:
            done(n)

    threading.Thread(target=run, daemon=True).start()


# ====================================================================== báo cáo lớp


def summarize(members: list[dict], results: list[dict], progress: list[dict]) -> list[dict]:
    """Mỗi học sinh một dòng: số lần làm, số lần đạt, điểm cao nhất từng môn, lần gần nhất, bài giảng đã học."""
    rows = []
    for m in members:
        mine = [r for r in results if r.get("hoc_vien") == m["id"]]
        best = {}
        for code in ("WORD", "EXCEL", "POWERPOINT"):
            scores = [r["score"] for r in mine if r.get("code") == code]
            best[code] = max(scores) if scores else None
        learned = [p for p in progress if p.get("hoc_vien") == m["id"]]
        rows.append({
            "id": m["id"], "username": m.get("username", ""), "ho_ten": m.get("ho_ten", ""),
            "khoa": bool(m.get("khoa")), "lan_lam": len(mine), "lan_dat": sum(bool(r.get("passed")) for r in mine),
            "best": best, "gan_nhat": max((r.get("lam_luc", "") for r in mine), default=""),
            "bai_xong": sum(bool(p.get("xong")) for p in learned), "bai_da_mo": len(learned),
        })
    return rows


def export_excel(path: Path, class_name: str, summary: list[dict], results: list[dict],
                 members: list[dict]) -> None:
    """Xuất bảng điểm lớp ra Excel: trang Tổng hợp + trang Chi tiết."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    names = {m["id"]: (m.get("username", ""), m.get("ho_ten", "")) for m in members}
    wb = Workbook()
    ws = wb.active
    ws.title = "Tong hop"
    ws.append([f"Lớp: {class_name}", "", f"Xuất lúc {datetime.now():%Y-%m-%d %H:%M}"])
    head = ["Tên đăng nhập", "Họ tên", "Số lần làm", "Số lần đạt", "Word cao nhất", "Excel cao nhất",
            "PowerPoint cao nhất", "Lần gần nhất", "Bài giảng đã học xong"]
    ws.append(head)
    for r in summary:
        ws.append([r["username"], r["ho_ten"], r["lan_lam"], r["lan_dat"], r["best"]["WORD"], r["best"]["EXCEL"],
                   r["best"]["POWERPOINT"], (r["gan_nhat"] or "")[:16].replace("T", " "), r["bai_xong"]])
    ws2 = wb.create_sheet("Chi tiet")
    ws2.append(["Tên đăng nhập", "Họ tên", "Thời gian", "Môn", "Bài thi", "Chế độ", "Điểm", "Đúng", "Tổng",
                "Đạt", "Thời lượng (phút)"])
    for r in results:
        user, name = names.get(r.get("hoc_vien"), ("", ""))
        ws2.append([user, name, (r.get("lam_luc") or "")[:16].replace("T", " "), r.get("code", ""), r.get("exam", ""),
                    r.get("mode", ""), r.get("score"), r.get("correct"), r.get("total"),
                    "Đạt" if r.get("passed") else "Chưa đạt", round((r.get("duration") or 0) / 60, 1)])
    bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2D4A2B")
    for sheet, row in ((ws, 2), (ws2, 1)):
        for cell in sheet[row]:
            cell.font, cell.fill = bold, fill
        for col in sheet.columns:
            width = max(len(str(c.value or "")) for c in col[row - 1:]) + 2
            sheet.column_dimensions[col[0].column_letter].width = min(max(width, 10), 50)
        sheet.freeze_panes = sheet.cell(row=row + 1, column=1)
    wb.save(path)
